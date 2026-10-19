"""Statystyki ogolne (rozliczenie): ekran z kosztami tylko dla admina, edycja linii recznych,
przelaczenie na skany, eksport Excel."""
import calendar
import io
from datetime import date, datetime

import openpyxl
from flask import Blueprint, abort, jsonify, make_response, render_template, request
from flask_login import current_user
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from logistat.czas import local_today
from logistat.extensions import db
from logistat.kategorie import (
    STAT_CATEGORIES, STAT_CATEGORY_LABELS, TOTAL_AMOUNT_BEZ, total_amount_wiersza)
from logistat.models import CostMapping, GeneralStat, nazwy_kategorii_pl
from logistat.rozliczenie import double_rate_amount_map, scan_category_totals, scan_coverage_map
from logistat.tabela_ilosci import ZRODLO_KATEGORII
from logistat.web.wspolne import acting_as_admin, admin_required, json_body, leader_required

bp = Blueprint('rozliczenie', __name__)


@bp.route('/general-stats/export')
@leader_required
def general_stats_export():
    # Lider dostaje ten sam plik bez kolumn kosztow — stawki widzi tylko admin.
    koszty = acting_as_admin()

    date_from_str = request.args.get('date_from', '')
    date_to_str = request.args.get('date_to', '')

    query = GeneralStat.query
    if date_from_str:
        try:
            query = query.filter(GeneralStat.loading_date >= datetime.strptime(date_from_str, '%Y-%m-%d').date())
        except ValueError:
            pass
    if date_to_str:
        try:
            query = query.filter(GeneralStat.loading_date <= datetime.strptime(date_to_str, '%Y-%m-%d').date())
        except ValueError:
            pass
    stats = query.order_by(GeneralStat.loading_date.asc()).all()

    dr_amounts = double_rate_amount_map()

    cost_mappings = CostMapping.query.all() if koszty else []
    rates_by_ym = {(cm.year, cm.month): cm.get_rates_data() for cm in cost_mappings}
    kol_na_kat = 2 if koszty else 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Statystyki ogólne'

    # Styles
    header_font = Font(bold=True, color='FFFFFF', size=10)
    header_fill_main = PatternFill('solid', fgColor='3B4A6B')
    header_fill_cat  = PatternFill('solid', fgColor='4F46E5')
    header_fill_dr   = PatternFill('solid', fgColor='C2410C')
    header_fill_tot  = PatternFill('solid', fgColor='166534')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Border(
        left=Side(style='thin', color='555555'),
        right=Side(style='thin', color='555555'),
        top=Side(style='thin', color='555555'),
        bottom=Side(style='thin', color='555555'),
    )
    dr_fill_row = PatternFill('solid', fgColor='FFF3E0')

    # Row 1: main headers
    base_headers = ['Loading date', 'Double Rate', 'Week', 'List-ID',
                    'Country of destination', 'Country ledger', 'Amounts', 'Total Amount']
    cat_cols_start = len(base_headers) + 1

    for ci, h in enumerate(base_headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = header_font
        cell.fill = header_fill_dr if h == 'Double Rate' else header_fill_main
        cell.alignment = center
        cell.border = thin
        ws.merge_cells(start_row=1, start_column=ci, end_row=2, end_column=ci)

    for i, cat in enumerate(STAT_CATEGORIES):
        col = cat_cols_start + i * kol_na_kat
        label = STAT_CATEGORY_LABELS[cat]
        c = ws.cell(row=1, column=col, value=label)
        c.font = header_font
        c.fill = header_fill_cat
        c.alignment = center
        c.border = thin
        if kol_na_kat > 1:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)

    total_col = cat_cols_start + len(STAT_CATEGORIES) * kol_na_kat
    if koszty:
        tc = ws.cell(row=1, column=total_col, value='Total cost (EUR)')
        tc.font = header_font
        tc.fill = header_fill_tot
        tc.alignment = center
        tc.border = thin
        ws.merge_cells(start_row=1, start_column=total_col, end_row=2, end_column=total_col)

    # Row 2: sub-headers for categories
    for i, cat in enumerate(STAT_CATEGORIES):
        col = cat_cols_start + i * kol_na_kat
        for offset, sub in enumerate(['Amount', 'Cost (EUR)'][:kol_na_kat]):
            c = ws.cell(row=2, column=col + offset, value=sub)
            c.font = Font(bold=True, color='FFFFFF', size=9)
            c.fill = header_fill_cat
            c.alignment = center
            c.border = thin

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18

    def write_data_row(ri, s, cd, dr_label, row_fill, amounts_val):
        """Write one data row (normal or yellow double-rate) at Excel row ri."""
        ym = (s.loading_date.year, s.loading_date.month)
        rates = rates_by_ym.get(ym, {})

        # Total Amount: wszystko poza twice/triple, w obu wierszach (TOTAL_AMOUNT_BEZ).
        total_amount = total_amount_wiersza(cd)
        total_cost = sum(cd.get(cat, {}).get('amount', 0) * rates.get(cat, 0.0)
                         for cat in STAT_CATEGORIES)

        row_data = [
            s.loading_date.strftime('%d.%m.%Y'),
            dr_label,
            s.week_number,
            s.list_id,
            s.country_of_destination or '',
            s.country_ledger,
            amounts_val,
            total_amount,
        ]

        for ci, val in enumerate(row_data, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.border = thin
            c.alignment = Alignment(horizontal='center' if ci in (1, 2, 3, 7, 8) else 'left')
            if row_fill:
                c.fill = row_fill
            if ci == 2 and dr_label:
                c.font = Font(bold=True, color='C2410C')

        for i, cat in enumerate(STAT_CATEGORIES):
            col = cat_cols_start + i * kol_na_kat
            amt = cd.get(cat, {}).get('amount', 0)
            cost = amt * rates.get(cat, 0.0)
            for offset, val in enumerate([amt, round(cost, 2)][:kol_na_kat]):
                c = ws.cell(row=ri, column=col + offset, value=val)
                c.border = thin
                c.alignment = Alignment(horizontal='right')
                if row_fill:
                    c.fill = row_fill

        if not koszty:
            return
        tc = ws.cell(row=ri, column=total_col, value=round(total_cost, 2))
        tc.border = thin
        tc.alignment = Alignment(horizontal='right')
        tc.font = Font(bold=True)
        if row_fill:
            tc.fill = row_fill

    # Data rows — normal row, then a yellow double-rate row when the line has
    # double-rate packages. The doubling is emergent: the same packages are
    # billed once in the normal row and again (×1) in the yellow row.
    ri = 3
    for s in stats:
        write_data_row(ri, s, s.get_category_data(), '', None, s.amounts)
        ri += 1

        dr_amount = dr_amounts.get((s.list_id, s.country_ledger, s.loading_date), 0)
        if dr_amount > 0:
            write_data_row(ri, s, s.get_double_rate_category_data(),
                           'DOUBLE RATE', dr_fill_row, dr_amount)
            ri += 1

    # Column widths
    col_widths = [12, 10, 6, 14, 22, 18, 9, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for i in range(len(STAT_CATEGORIES) * kol_na_kat):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cat_cols_start + i)].width = 9
    if koszty:
        ws.column_dimensions[openpyxl.utils.get_column_letter(total_col)].width = 14

    ws.freeze_panes = 'A3'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    fname = f"statystyki_{date_from_str or 'all'}_{date_to_str or 'all'}.xlsx"
    response = make_response(output.read())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename="{fname}"'
    return response


@bp.route('/general-stats')
@leader_required
def general_stats_page():
    # Lider: podglad bez kosztow i bez edycji (API zapisu zostaje @admin_required).
    # Stawek nie wysylamy mu nawet do JS — bylyby w zrodle strony.
    admin = acting_as_admin()
    # If parameters not provided, use first and last day of current month
    today = local_today()
    default_from = today.replace(day=1).strftime('%Y-%m-%d')
    # Calculate last day of month
    last_day = calendar.monthrange(today.year, today.month)[1]
    default_to = today.replace(day=last_day).strftime('%Y-%m-%d')

    date_from_str = request.args.get('date_from', default_from)
    date_to_str = request.args.get('date_to', default_to)
    
    query = GeneralStat.query
    
    date_from = None
    date_to = None
    
    if date_from_str:
        try:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date()
            query = query.filter(GeneralStat.loading_date >= date_from)
        except ValueError:
            pass
            
    if date_to_str:
        try:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date()
            query = query.filter(GeneralStat.loading_date <= date_to)
        except ValueError:
            pass

    stats = query.order_by(GeneralStat.loading_date.desc()).all()

    # Ilość double-rate (suma Stückzahl paczek oznaczonych) per linia → żółty wiersz
    dr_amounts = double_rate_amount_map()
    pokrycie = scan_coverage_map()
    for s in stats:
        s.dr_amount = dr_amounts.get((s.list_id, s.country_ledger, s.loading_date), 0)
        # Linia zasilana skanami czyta w trakcie zmiany mniej niz finalnie —
        # licznik pokrycia pilnuje, by nikt nie wzial jej za gotowa.
        s.scanned_cartons, s.total_cartons = pokrycie.get(
            (s.list_id, s.country_ledger, s.loading_date), (0, 0))

    # Przekazanie do szablonu wszystkich mapowań kosztów aby JS miał do nich dostęp przy edycji
    cost_mappings = CostMapping.query.all() if admin else []
    rates_by_ym = {f"{cm.year}-{cm.month:02d}": cm.get_rates_data() for cm in cost_mappings}

    return render_template('general_stats.html',
                           stats=stats,
                           category_labels=STAT_CATEGORY_LABELS,
                           category_labels_pl=nazwy_kategorii_pl(),
                           total_amount_bez=list(TOTAL_AMOUNT_BEZ),
                           zrodlo_kategorii=ZRODLO_KATEGORII,
                           categories=STAT_CATEGORIES,
                           date_from=date_from_str,
                           date_to=date_to_str,
                           rates_by_ym=rates_by_ym,
                           pokaz_koszty=admin,
                           edycja=admin)


@bp.route('/api/general-stats', methods=['GET'])
@admin_required
def api_general_stats():
    """Lista linii rozliczeniowych. Opcjonalne `date_from` / `date_to`
    (YYYY-MM-DD) — bez nich zwraca wszystko, jak dotad."""
    query = GeneralStat.query
    for param, warunek in (('date_from', GeneralStat.loading_date.__ge__),
                           ('date_to', GeneralStat.loading_date.__le__)):
        wartosc = request.args.get(param, '')
        if wartosc:
            try:
                query = query.filter(warunek(date.fromisoformat(wartosc)))
            except ValueError:
                return jsonify({'error': f'Nieprawidłowa data w {param}.'}), 400
    stats = query.order_by(GeneralStat.loading_date.desc()).all()
    # Pokrycie liczymy raz dla calej listy (jedno zapytanie), nie per wiersz.
    pokrycie = scan_coverage_map()
    wynik = []
    for s in stats:
        d = s.to_dict()
        zesk, wsz = pokrycie.get((s.list_id, s.country_ledger, s.loading_date), (0, 0))
        d['scanned_cartons'] = zesk
        d['total_cartons'] = wsz
        wynik.append(d)
    return jsonify(wynik), 200


@bp.route('/api/general-stats/<int:stat_id>', methods=['PUT'])
@admin_required
def api_general_stat_update(stat_id):
    stat = db.get_or_404(GeneralStat, stat_id)
    data = json_body()

    if 'category_data' in data:
        if (stat.category_source or 'manual') == 'scan':
            abort(400, 'Kategorie tej linii pochodzą ze skanów paczek — '
                       'popraw ilości przy paczce (Paczki → ✎ Ilości).')
        stat.set_category_data(data['category_data'])
    if 'double_rate_category_data' in data:
        stat.set_double_rate_category_data(data['double_rate_category_data'])
    stat.updated_at = datetime.utcnow()
    stat.updated_by = current_user.id

    db.session.commit()
    return jsonify(stat.to_dict()), 200


@bp.route('/api/general-stats/<int:stat_id>/use-scan', methods=['POST'])
@admin_required
def api_general_stat_use_scan(stat_id):
    """Swiadome przelaczenie linii z ilosci recznych na sume skanow.

    Linia z recznymi ilosciami nie przelacza sie sama (patrz recompute_general_stat),
    zeby pierwszy zeskanowany karton nie zastapil calego rozliczenia. Tu admin
    potwierdza zamiane — po niej `category_data` jest suma skanow tej grupy.
    """
    stat = db.get_or_404(GeneralStat, stat_id)
    stat.category_source = 'scan'
    stat.set_category_data(
        scan_category_totals(stat.list_id, stat.country_ledger, stat.loading_date))
    stat.updated_at = datetime.utcnow()
    stat.updated_by = current_user.id
    db.session.commit()
    return jsonify(stat.to_dict()), 200
