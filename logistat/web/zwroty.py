"""Zwroty (kafelek od 2026-10-24).

- „Zwroty Paczek” (2026-10-26): tabela jak w arkuszu operacji Return_<miesiac>,
  filtry, eksport do Excela.
- „Invoice specification” i „Sorting Data” (2026-10-26): raporty miesieczne /
  roczne w ukladzie arkuszy operacji; ilosci wpisuje lider (autozapis), ceny
  ustawia admin w Ustawieniach (miesiacami), listy wierszy Sorting Data tez admin.
  Lider nie widzi cen ani Total charge — ani na ekranie, ani w Excelu.
- „Przyjęcie zwrotu” (2026-10-27, kafelek Magazyn): pracownik przechodzi ekran
  po ekranie, jeden zapis = jeden wiersz Zwrotow Paczek. Napisy ekranow
  i przyciski (brandy, kurierzy, Verpackungsvorschrift) ustawia admin.
  Bez osob przydzielonych do stanowiska przyjecie jest zablokowane (409).
- „Przydział zwrotów” i „Dashboard zwrotów” (2026-10-27): lider przydziela osoby
  z czynnosci Zwroty do stanowisk / Supportu; statystyki dni miesiaca.
"""
import logging
from datetime import date, timedelta

from flask import Blueprint, abort, jsonify, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.exc import IntegrityError

from logistat.czas import local_today, teraz_utc
from logistat.extensions import db
from logistat.faktura_zwrotow import (
    CZYNNOSCI_AUTO, CZYNNOSCI_FAKTURY, KLUCZ_POS, KLUCZE_FAKTURY, MAX_ILOSC_ZWROTOW,
    MIESIACE_EN, NAZWA_LICENCJI_MAX, POLA_SORTING, RODZAJE_LICENCJI)
from logistat.models import (
    LicencjaZwrotow, Shift, SortingDataWpis, User, ZwrotPaczki, get_or_create_shift)
from logistat.przydzial_zwrotow import (
    dashboard_zwrotow, obsada_stanowiska, przydziel, stan_przydzialu, stanowiska_zwrotow,
    zamknij_otwarty, zmiana_edytowalna)
from logistat.przyjecie_zwrotu import (
    BRAK_POS, BRAND, ETYKIETA_MAX, ETYKIETY_DOMYSLNE, KOLUMNA_ETYKIETY, KURIER, LICZNIKI,
    KOMUNIKAT_BRAK_OBSADY, MAX_GODZIN_PRZYDZIALU, MAX_LICZNIK, MIESIACE_PL, OPAKOWANIE, PAKETNUMMER_MAX, RODZAJE_PRZYJECIA, SUPPORT,
    TOKEN_MAX, pos_nummer)
from logistat.web.excel import odpowiedz_xlsx
from logistat.web.wspolne import (
    acting_as_admin, admin_required, json_body, leader_required, magazyn_required, parse_date,
    parse_shift_number, require_int)
from logistat.zwroty_dane import (
    brandy_zwrotow, ceny_miesiaca, etykiety_przyjecia, granice_miesiaca, invoice_specification,
    kurierzy_raportu, listy_przyjecia, naglowki_zwrotow, opcje_przyjecia, sorting_data,
    wiersze_zwrotow, zapisz_ceny, zapisz_etykiety, zapisz_ilosc_miesiaca, zapisz_sorting,
    zapisz_zwrot)

# Listy edytowane w Ustawieniach — Sorting Data i przyciski Przyjecia zwrotu.
RODZAJE_LIST = {**RODZAJE_LICENCJI, **RODZAJE_PRZYJECIA}

bp = Blueprint('zwroty', __name__)
log = logging.getLogger(__name__)

# Szerokosci kolumn z arkusza operacji (Return_September_2026): przed
# kurierami, kolumna kuriera (jak DPD / UPS), po kurierach.
SZEROKOSCI_PRZED = (22.1, 10.7, 13.7, 17.9)
SZEROKOSC_KURIERA = 9.3
SZEROKOSCI_PO = (12.7, 11.6, 26.3, 19.3, 28.6, 20.9, 25.6, 13.7, 11.7)


@bp.route('/zwroty')
@leader_required
def zwroty():
    # Pusta pozycja usunieta z menu (operacje 2026-10-26); adres zostaje
    # (zakladki, mapa tras) i prowadzi na ekran startowy kafelka.
    return redirect(url_for('zwroty.zwroty_paczek'))


def _dzien(wartosc):
    """Data z filtra; pusta albo bledna → None (strona Jinja — bez 400)."""
    try:
        return date.fromisoformat(wartosc) if wartosc else None
    except ValueError:
        return None


def _filtry():
    """Filtry z query stringu. Domyslnie tylko dzisiejszy dzien (operacja)."""
    dzis = local_today()
    od_arg = _dzien(request.args.get('date_from', ''))
    do_arg = _dzien(request.args.get('date_to', ''))
    # Jedna granica = jeden dzien; zadnej = dzis.
    od = od_arg or do_arg or dzis
    do = do_arg or od
    if od > do:
        od, do = do, od
    return {
        'od': od, 'do': do,
        'brand': request.args.get('brand', '').strip() or None,
        'paketnummer': request.args.get('paketnummer', '').strip() or None,
    }


@bp.route('/zwroty/paczki')
@leader_required
def zwroty_paczek():
    f = _filtry()
    kurierzy = kurierzy_raportu(f['od'], f['do'])
    return render_template('zwroty_paczek.html', naglowki=naglowki_zwrotow(kurierzy),
                           wiersze=wiersze_zwrotow(**f, kurierzy=kurierzy),
                           brandy=brandy_zwrotow(), filtry=f, dzis=local_today(),
                           listy=listy_przyjecia(), liczniki=LICZNIKI)


@bp.route('/zwroty/paczki/export')
@leader_required
def zwroty_paczek_export():
    f = _filtry()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Returns'
    kurierzy = kurierzy_raportu(f['od'], f['do'])
    ws.append(list(naglowki_zwrotow(kurierzy)))
    for komorka in ws[1]:
        komorka.alignment = Alignment(horizontal='center')
    for wiersz in wiersze_zwrotow(**f, kurierzy=kurierzy):
        ws.append(list(wiersz))
        nr = ws.max_row
        # Paketnummer zaczyna sie od „%” i ma wiodace zera — zawsze tekst.
        ws.cell(nr, 1).number_format = '@'
        if isinstance(ws.cell(nr, 3).value, date):
            ws.cell(nr, 3).number_format = 'DD.MM.YYYY'
    szerokosci = SZEROKOSCI_PRZED + (SZEROKOSC_KURIERA,) * len(kurierzy) + SZEROKOSCI_PO
    for i, szer in enumerate(szerokosci, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = szer
    ws.auto_filter.ref = ws.dimensions
    nazwa = (f'Returns_{f["od"]:%Y-%m-%d}.xlsx' if f['od'] == f['do']
             else f'Returns_{f["od"]:%Y-%m-%d}_{f["do"]:%Y-%m-%d}.xlsx')
    return odpowiedz_xlsx(wb, nazwa)


# ── wspolne: rok / miesiac / liczby ──────────────────────────────────────────

def _okres_z_adresu(z_miesiacem=True):
    """Rok (i miesiac) z query stringu strony; bledne / brak → biezacy (bez 400)."""
    dzis = local_today()
    rok = request.args.get('rok', type=int)
    if rok is None or not 2000 <= rok <= 2100:
        rok = dzis.year
    if not z_miesiacem:
        return rok, None
    miesiac = request.args.get('miesiac', type=int)
    if miesiac is None or not 1 <= miesiac <= 12:
        miesiac = dzis.month
    return rok, miesiac


def _rok_miesiac(rok, miesiac):
    """Walidacja okresu z API (rozliczenie — nic na slowo)."""
    for wartosc, pole, od, do in ((rok, 'rok', 2000, 2100), (miesiac, 'miesiac', 1, 12)):
        if isinstance(wartosc, bool) or not isinstance(wartosc, int) or not od <= wartosc <= do:
            abort(400, f'Nieprawidłowe pole „{pole}”.')
    return rok, miesiac


def _ilosc(wartosc, pole='ilosc'):
    """Wpisana liczba: None / '' = puste pole, inaczej calkowita 0 … MAX."""
    if wartosc is None or wartosc == '':
        return None
    if isinstance(wartosc, float) and wartosc.is_integer():
        wartosc = int(wartosc)
    if isinstance(wartosc, bool) or not isinstance(wartosc, int):
        abort(400, f'Pole „{pole}” musi być liczbą całkowitą.')
    if not 0 <= wartosc < MAX_ILOSC_ZWROTOW:
        abort(400, f'Pole „{pole}” musi być liczbą od 0 do {MAX_ILOSC_ZWROTOW - 1}.')
    return wartosc


# ── Invoice specification ────────────────────────────────────────────────────

@bp.route('/zwroty/invoice-specification')
@leader_required
def invoice_specification_view():
    rok, miesiac = _okres_z_adresu()
    dane = invoice_specification(rok, miesiac)
    pokaz_ceny = acting_as_admin()
    if not pokaz_ceny:
        # Lider: ceny i kwoty nie trafiaja nawet do zrodla strony.
        for w in dane['wiersze']:
            w['cena'] = w['charge'] = None
        dane['total'] = None
        dane['brak_cen'] = []
    return render_template('invoice_specification.html', dane=dane, pokaz_ceny=pokaz_ceny,
                           miesiace=MIESIACE_EN)


@bp.route('/api/zwroty/invoice-specification/ilosc', methods=['PUT'])
@leader_required
def api_invoice_ilosc():
    d = json_body()
    rok, miesiac = _rok_miesiac(d.get('rok'), d.get('miesiac'))
    klucz = d.get('klucz')
    if klucz not in KLUCZE_FAKTURY:
        abort(400, 'Nieznana czynność.')
    if klucz in CZYNNOSCI_AUTO:
        abort(400, 'Sorting returns liczy system z raportu Zwroty Paczek — nie wpisuje się go ręcznie.')
    zapisz_ilosc_miesiaca(rok, miesiac, klucz, _ilosc(d.get('ilosc')), current_user.id)
    return jsonify({'ok': True})


# Kolory z arkusza operacji (motyw Office: akcent 5 / 6 / ciemny 2, rozjasnione).
_TLO_TYTULU = PatternFill('solid', fgColor='BDD7EE')
_TLO_NAGLOWKA = PatternFill('solid', fgColor='C6E0B4')
_TLO_OKRESU = PatternFill('solid', fgColor='D6DCE4')
_FORMAT_EUR = '#,##0.00\\ [$EUR];\\-#,##0.00\\ [$EUR]'
_SRODEK = Alignment(horizontal='center')


@bp.route('/zwroty/invoice-specification/export')
@leader_required
def invoice_specification_export():
    rok, miesiac = _okres_z_adresu()
    dane = invoice_specification(rok, miesiac)
    pokaz_ceny = acting_as_admin()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoice specification'
    ostatnia = 6 if pokaz_ceny else 4      # kolumna F (Total charge) albo D (UoM)
    for kol, szer in zip('ABCDEF', (8.9, 34.7, 15.0, 8.1, 9.3, 13.3)):
        ws.column_dimensions[kol].width = szer

    ws['B3'] = 'Warehousing Activities'
    ws['B3'].font = Font(bold=True, size=14)
    ws['B4'] = 'Returns logistics service'
    ws['B4'].font = Font(bold=True, size=12)
    for kol in range(2, ostatnia + 1):
        ws.cell(3, kol).fill = _TLO_TYTULU
        ws.cell(4, kol).border = Border(bottom=Side(style='thin'))
    for nr, (etykieta, wartosc) in ((6, ('Year', rok)), (7, ('Month', miesiac))):
        ws.cell(nr, 2, etykieta).fill = _TLO_OKRESU
        c = ws.cell(nr, 3, wartosc)
        c.fill = _TLO_OKRESU
        c.alignment = Alignment(horizontal='right')

    naglowki = ['Activity', 'Total quantity', 'UoM'] + (['Price', 'Total charge'] if pokaz_ceny else [])
    for i, tekst in enumerate(naglowki, start=2):
        c = ws.cell(9, i, tekst)
        c.fill = _TLO_NAGLOWKA
        if i > 2:
            c.alignment = _SRODEK

    nr = 11
    for w in dane['wiersze']:
        ws.cell(nr, 2, w['activity'])
        ws.cell(nr, 3, w['ilosc']).alignment = _SRODEK
        ws.cell(nr, 4, w['uom']).alignment = _SRODEK
        if pokaz_ceny:
            c = ws.cell(nr, 5, w['cena'])
            c.alignment, c.number_format = _SRODEK, _FORMAT_EUR
            # Formula jak w arkuszu — plik dalej liczy po poprawce w Excelu.
            c = ws.cell(nr, 6, f'=E{nr}*C{nr}')
            c.number_format = _FORMAT_EUR
        nr += 1
    ostatni_wiersz = nr - 1
    nr += 1
    if pokaz_ceny:
        for kol in range(2, 7):
            ws.cell(nr, kol).fill = _TLO_NAGLOWKA
        ws.cell(nr, 2, 'Total charge')
        c = ws.cell(nr, 6, f'=SUM(F11:F{ostatni_wiersz})')
        c.number_format = _FORMAT_EUR
    return odpowiedz_xlsx(wb, f'Invoice_specification_{rok}-{miesiac:02d}.xlsx')


# ── Sorting Data ─────────────────────────────────────────────────────────────

@bp.route('/zwroty/sorting-data')
@leader_required
def sorting_data_view():
    rok, _ = _okres_z_adresu(z_miesiacem=False)
    return render_template('sorting_data.html', dane=sorting_data(rok), miesiace=MIESIACE_EN)


@bp.route('/api/zwroty/sorting-data', methods=['PUT'])
@leader_required
def api_sorting_data():
    """Jedna komorka: {rok, miesiac, pole: cartons|pallets, licencja_id, wartosc}
    albo POS processed: {rok, miesiac, pole: 'pos_processed', wartosc}."""
    d = json_body()
    rok, miesiac = _rok_miesiac(d.get('rok'), d.get('miesiac'))
    pole = d.get('pole')
    wartosc = _ilosc(d.get('wartosc'), 'wartosc')
    if pole == KLUCZ_POS:
        zapisz_ilosc_miesiaca(rok, miesiac, KLUCZ_POS, wartosc, current_user.id)
        return jsonify({'ok': True})
    if pole not in POLA_SORTING:
        abort(400, 'Nieznane pole.')
    lic_id = d.get('licencja_id')
    lic = db.session.get(LicencjaZwrotow, lic_id) if isinstance(lic_id, int) and not isinstance(lic_id, bool) else None
    if lic is None:
        abort(400, 'Nieznana licencja / wysyłka.')
    zapisz_sorting(lic.id, rok, miesiac, pole, wartosc, current_user.id)
    return jsonify({'ok': True})


# Kolory miesiecy z arkusza (na zmiane ciemniejszy / jasniejszy szary).
_TLA_MIESIECY = (PatternFill('solid', fgColor='BFBFBF'), PatternFill('solid', fgColor='F2F2F2'))
_TLO_SUMY = PatternFill('solid', fgColor='BFBFBF')
_CIENKA = Side(style='thin')
_RAMKA = Border(left=_CIENKA, right=_CIENKA, top=_CIENKA, bottom=_CIENKA)


@bp.route('/zwroty/sorting-data/export')
@leader_required
def sorting_data_export():
    rok, _ = _okres_z_adresu(z_miesiacem=False)
    dane = sorting_data(rok)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sorting Data'
    ws.column_dimensions['A'].width = 24
    for m in range(1, 13):
        kol = 2 * m              # B, D, F … = Cartons; C, E, G … = Pallets
        tlo = _TLA_MIESIECY[(m - 1) % 2]
        ws.merge_cells(start_row=2, start_column=kol, end_row=2, end_column=kol + 1)
        c = ws.cell(2, kol, MIESIACE_EN[m - 1])
        c.font, c.alignment = Font(bold=True), _SRODEK
        for k, tekst in ((kol, 'Cartons'), (kol + 1, 'Pallets')):
            for r in (2, 3):
                ws.cell(r, k).fill, ws.cell(r, k).border = tlo, _RAMKA
            ws.cell(3, k, tekst).alignment = _SRODEK
            ws.column_dimensions[get_column_letter(k)].width = 9

    def wiersz(nr, nazwa, mies=None, tlo=None, formula_sumy=None, tylko_cartons=None):
        ws.cell(nr, 1, nazwa).border = Border(left=_CIENKA)
        for m in range(1, 13):
            for i, pole in enumerate(POLA_SORTING):
                kol = 2 * m + i
                if formula_sumy:
                    wartosc = formula_sumy(get_column_letter(kol))
                elif tylko_cartons is not None:
                    wartosc = tylko_cartons.get(m) if pole == 'cartons' else None
                else:
                    wartosc = mies[m][pole]
                c = ws.cell(nr, kol, wartosc)
                c.alignment, c.border = _SRODEK, _RAMKA
                if tlo:
                    c.fill = tlo
        if tlo:
            ws.cell(nr, 1).fill = tlo

    nr = 4
    for w in dane['licencje']:
        wiersz(nr, w['nazwa'], w['mies'])
        nr += 1
    pierwszy, ostatni = 4, nr - 1
    # Suma jako formula (jak w arkuszu) — poprawka w Excelu przelicza „Total”.
    wiersz(nr, 'Total boxes built', tlo=_TLO_SUMY,
           formula_sumy=(lambda k: f'=SUM({k}{pierwszy}:{k}{ostatni})') if ostatni >= pierwszy
           else (lambda k: 0))
    nr += 1
    for w in dane['wysylki']:
        wiersz(nr, w['nazwa'], w['mies'])
        nr += 1
    wiersz(nr, 'POS processed', tylko_cartons=dane['pos'])
    return odpowiedz_xlsx(wb, f'Sorting_Data_{rok}.xlsx')


# ── Ustawienia (admin): listy Sorting Data i cennik Invoice specification ────

@bp.route('/admin/zwroty/licencje')
@admin_required
def admin_licencje():
    lista = LicencjaZwrotow.query.order_by(LicencjaZwrotow.kolejnosc, LicencjaZwrotow.id).all()
    uzywane = {i for (i,) in db.session.query(SortingDataWpis.licencja_id).distinct()}
    return render_template('admin_licencje_zwrotow.html', rodzaje=RODZAJE_LICENCJI,
                           licencje={r: [lic for lic in lista if lic.rodzaj == r] for r in RODZAJE_LICENCJI},
                           uzywane=uzywane, max_nazwy=NAZWA_LICENCJI_MAX)


# ── Przyjecie zwrotu (kafelek Magazyn) ───────────────────────────────────────

@bp.route('/zwroty/przyjecie')
@magazyn_required
def przyjecie_zwrotu():
    obsada = _obsada_nazwy(current_user, teraz_utc())
    return render_template('przyjecie_zwrotu.html', etykiety=etykiety_przyjecia(),
                           listy=listy_przyjecia(), liczniki=LICZNIKI, brak_pos=BRAK_POS,
                           max_licznik=MAX_LICZNIK, obsada=obsada,
                           brak_obsady=KOMUNIKAT_BRAK_OBSADY)


def _obsada_nazwy(konto, chwila):
    """Imiona osob na stanowisku tego konta; [] = nikt (albo to nie stanowisko)."""
    if konto.role != 'stanowisko':
        return []
    ids = obsada_stanowiska(konto.id, chwila)
    if not ids:
        return []
    return [u.display_name for u in User.query.filter(User.id.in_(ids)).order_by(User.display_name)]


def _z_listy(d, pole, rodzaj, nazwa_pola, obecna=None):
    """Wartosc z listy przyciskow: aktywna pozycja albo — przy poprawce — ta,
    ktora zwrot juz ma (wylaczony brand nie blokuje poprawienia reszty)."""
    wartosc = d.get(pole)
    if not isinstance(wartosc, str) or (wartosc not in opcje_przyjecia(rodzaj) and wartosc != obecna):
        # Admin mogl wlasnie wylaczyc pozycje — ekran pokazuje stara liste.
        abort(400, f'Nieznana wartość „{nazwa_pola}” — odśwież stronę.')
    return wartosc


def _dane_zwrotu(d, obecny=None):
    """Pola zwrotu z body (przyjecie na stanowisku i poprawka lidera — te same
    zasady). Date / Monat / Woche nie przychodza nigdy: to data przyjecia."""
    paket = d.get('paketnummer')
    paket = paket.strip() if isinstance(paket, str) else ''
    if not paket:
        abort(400, 'Zeskanuj numer paczki.')
    if len(paket) > PAKETNUMMER_MAX:
        abort(400, f'Numer paczki może mieć najwyżej {PAKETNUMMER_MAX} znaków.')
    pos = pos_nummer(d.get('pos_nummer'))
    if pos is None:
        abort(400, 'Numer POS: tylko litery, cyfry i - _ / . (albo przycisk „brak numeru”).')
    dane = {
        'paketnummer': paket, 'pos_nummer': pos,
        'brand': _z_listy(d, 'brand', BRAND, 'Brand', obecny and obecny.brand),
        'kurier': _z_listy(d, 'kurier', KURIER, 'kurier', obecny and obecny.kurier),
        'verpackungsvorschrift': _z_listy(d, 'verpackungsvorschrift', OPAKOWANIE,
                                          'Verpackungsvorschrift',
                                          obecny and obecny.verpackungsvorschrift),
    }
    for pole, naglowek, minimum, _ in LICZNIKI:
        wartosc = d.get(pole)
        if isinstance(wartosc, bool) or not isinstance(wartosc, int) \
                or not minimum <= wartosc < MAX_LICZNIK:
            abort(400, f'{naglowek}: liczba całkowita od {minimum} do {MAX_LICZNIK - 1}.')
        dane[pole] = wartosc
    return dane


@bp.route('/api/zwroty/przyjecie', methods=['POST'])
@magazyn_required
def api_przyjecie_zwrotu():
    """Jeden przyjety zwrot. Date / Monat / Woche ustawia serwer (data
    przyjecia), nigdy przegladarka. Ten sam `token` drugi raz → 200 bez nowego
    wiersza (`powtorzony`), zeby podwojny klik nie liczyl sie na fakturze 2×."""
    d = json_body()
    token = d.get('token')
    if not isinstance(token, str) or not 8 <= len(token) <= TOKEN_MAX:
        abort(400, 'Brak identyfikatora zapisu — odśwież stronę.')
    dane = _dane_zwrotu(d)
    # Kolejnosc: walidacja → ten sam token (ponowienie juz zapisanego zwrotu
    # przechodzi, nawet gdy lider w miedzyczasie zdjal obsade) → obsada → zapis.
    juz = ZwrotPaczki.query.filter_by(token=token).first()
    if juz is not None:
        return jsonify({'ok': True, 'id': juz.id, 'powtorzony': True,
                        'paketnummer': juz.paketnummer}), 200
    teraz = teraz_utc()
    obsada = _obsada_nazwy(current_user, teraz)
    if not obsada:
        return jsonify({'error': KOMUNIKAT_BRAK_OBSADY, 'brak_obsady': True}), 409
    zwrot, nowy = zapisz_zwrot(dane, token, current_user.id, teraz)
    return jsonify({'ok': True, 'id': zwrot.id, 'powtorzony': not nowy,
                    'paketnummer': zwrot.paketnummer, 'obsada': obsada}), 201 if nowy else 200


@bp.route('/api/zwroty/przyjecie/<int:zwrot_id>', methods=['PUT'])
@leader_required
def api_zwrot_popraw(zwrot_id):
    """Poprawka lidera (2026-10-27, operacja): te same pola co przy przyjeciu.
    Data przyjecia i stanowisko zostaja — od nich zalezy podzial na osoby."""
    zwrot = db.get_or_404(ZwrotPaczki, zwrot_id)
    for pole, wartosc in _dane_zwrotu(json_body(), zwrot).items():
        setattr(zwrot, pole, wartosc)
    zwrot.zmienil_id = current_user.id
    zwrot.zmieniono_at = teraz_utc()
    db.session.commit()
    return jsonify({'ok': True, 'id': zwrot.id})


@bp.route('/api/zwroty/przyjecie/<int:zwrot_id>', methods=['DELETE'])
@leader_required
def api_zwrot_usun(zwrot_id):
    """Usuniecie blednie przyjetego zwrotu (np. paczka zeskanowana dwa razy) —
    znika z raportu, Sorting returns i dashboardu."""
    zwrot = db.get_or_404(ZwrotPaczki, zwrot_id)
    log.info('Zwrot %s (%s, przyjety %s) usuniety przez %s', zwrot.id, zwrot.paketnummer,
             zwrot.przyjeto_at, current_user.username)
    db.session.delete(zwrot)
    db.session.commit()
    return jsonify({'ok': True})


# ── Przydzial zwrotow (lider) ────────────────────────────────────────────────

def _zmiana_przydzialu(dzien, numer):
    """Przydzial dziala „od teraz”, wiec tylko biezace zmiany: dzis albo wczoraj
    (druga zmiana pracuje po polnocy). Data i godzina z serwera."""
    if not zmiana_edytowalna(dzien):
        abort(400, 'Przydział zwrotów działa tylko dla bieżącej zmiany (dziś albo wczoraj).')
    return get_or_create_shift(dzien, numer)


@bp.route('/zwroty/przydzial')
@leader_required
def przydzial_zwrotow():
    dzien = _dzien(request.args.get('date', '')) or local_today()
    numer = request.args.get('shift', type=int)
    return render_template('przydzial_zwrotow.html', dzien=dzien,
                           zmiana=numer if numer in (1, 2) else 1, max_godzin=MAX_GODZIN_PRZYDZIALU)


@bp.route('/api/zwroty/przydzial', methods=['GET'])
@leader_required
def api_przydzial_zwrotow():
    dzien = parse_date(request.args.get('date'))
    numer = parse_shift_number(request.args.get('shift'))
    shift = get_or_create_shift(dzien, numer)
    return jsonify(stan_przydzialu(shift))


@bp.route('/api/zwroty/przydzial', methods=['PUT'])
@leader_required
def api_przydzial_zwrotow_zapisz():
    """{date, shift, user_id, miejsce: id stanowiska | 'support' | null}.
    Przeniesienie liczy sie od chwili zapisu; null = zejscie ze stanowiska."""
    d = json_body()
    dzien = parse_date(d.get('date'))
    numer = parse_shift_number(d.get('shift'))
    user_id = require_int(d.get('user_id'), 'user_id')
    miejsce = d.get('miejsce')
    shift = _zmiana_przydzialu(dzien, numer)
    # Dwoch liderow na jednej zmianie — kolejno, nie na raz.
    shift = Shift.query.filter_by(id=shift.id).with_for_update().one()
    teraz = teraz_utc()
    if miejsce is None:
        zamknij_otwarty(user_id, teraz)
    else:
        stan = stan_przydzialu(shift)
        if user_id not in {o['user_id'] for o in stan['osoby'] if not o['poza_zwrotami']}:
            abort(400, f'Ta osoba nie jest przydzielona do czynności „{stan["czynnosc"]}” '
                       'na tej zmianie — najpierw Przydzielanie.')
        if miejsce == SUPPORT:
            przydziel(shift, user_id, None, True, teraz, current_user.id)
        else:
            if isinstance(miejsce, bool) or not isinstance(miejsce, int) \
                    or miejsce not in {s.id for s in stanowiska_zwrotow()}:
                abort(400, 'Nieznane stanowisko.')
            przydziel(shift, user_id, miejsce, False, teraz, current_user.id)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        abort(409, 'Ktoś właśnie zmienił przydział tej osoby — odśwież ekran.')
    return jsonify(stan_przydzialu(shift))


# ── Dashboard zwrotow ────────────────────────────────────────────────────────

@bp.route('/zwroty/dashboard')
@leader_required
def dashboard_zwrotow_view():
    rok, miesiac = _okres_z_adresu()
    od, do = granice_miesiaca(rok, miesiac)
    dni = [od + timedelta(days=i) for i in range((do - od).days + 1)]
    return render_template('dashboard_zwrotow.html', d=dashboard_zwrotow(dni, teraz_utc()),
                           rok=rok, miesiac=miesiac, dzis=local_today(), MIESIACE_PL=MIESIACE_PL)


@bp.route('/admin/zwroty/przyjecie')
@admin_required
def admin_przyjecie():
    lista = LicencjaZwrotow.query.filter(LicencjaZwrotow.rodzaj.in_(RODZAJE_PRZYJECIA))\
        .order_by(LicencjaZwrotow.kolejnosc, LicencjaZwrotow.id).all()
    return render_template('admin_przyjecie_zwrotu.html', etykiety=etykiety_przyjecia(),
                           domyslne=ETYKIETY_DOMYSLNE, kolumny=KOLUMNA_ETYKIETY,
                           max_etykiety=ETYKIETA_MAX, rodzaje=RODZAJE_PRZYJECIA,
                           licencje={r: [x for x in lista if x.rodzaj == r] for r in RODZAJE_PRZYJECIA},
                           uzywane=set(), max_nazwy=NAZWA_LICENCJI_MAX)


@bp.route('/api/zwroty/etykiety', methods=['PUT'])
@admin_required
def api_etykiety_przyjecia():
    """{etykiety: {klucz: napis}} — pusty napis = powrot do domyslnego."""
    etykiety = json_body().get('etykiety')
    if not isinstance(etykiety, dict):
        abort(400, 'Pole „etykiety” musi być obiektem {klucz: napis}.')
    czyste = {}
    for klucz, napis in etykiety.items():
        if klucz not in ETYKIETY_DOMYSLNE:
            abort(400, f'Nieznany ekran: {klucz}.')
        if napis is not None and not isinstance(napis, str):
            abort(400, 'Napis musi być tekstem.')
        napis = ' '.join((napis or '').split())
        if len(napis) > ETYKIETA_MAX:
            abort(400, f'Napis może mieć najwyżej {ETYKIETA_MAX} znaków.')
        czyste[klucz] = napis
    zapisz_etykiety(czyste)
    return jsonify({'ok': True, 'etykiety': etykiety_przyjecia()})


def _nazwa_licencji(wartosc):
    nazwa = (wartosc or '').strip() if isinstance(wartosc, str) else ''
    if not nazwa:
        abort(400, 'Nazwa nie może być pusta.')
    if len(nazwa) > NAZWA_LICENCJI_MAX:
        abort(400, f'Nazwa może mieć najwyżej {NAZWA_LICENCJI_MAX} znaków.')
    return nazwa


def _zapisz_licencje():
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        abort(409, 'Taka nazwa już jest na tej liście.')


@bp.route('/api/zwroty/licencje', methods=['POST'])
@admin_required
def api_licencja_dodaj():
    d = json_body()
    rodzaj = d.get('rodzaj')
    if rodzaj not in RODZAJE_LIST:
        abort(400, 'Nieznana lista.')
    nazwa = _nazwa_licencji(d.get('nazwa'))
    ostatnia = db.session.query(db.func.max(LicencjaZwrotow.kolejnosc))\
        .filter(LicencjaZwrotow.rodzaj == rodzaj).scalar()
    lic = LicencjaZwrotow(rodzaj=rodzaj, nazwa=nazwa, kolejnosc=(ostatnia or 0) + 1)
    db.session.add(lic)
    _zapisz_licencje()
    return jsonify(lic.to_dict()), 201


@bp.route('/api/zwroty/licencje/<int:lic_id>', methods=['PUT'])
@admin_required
def api_licencja_zmien(lic_id):
    """{nazwa} — zmiana nazwy; {aktywna} — wlaczenie / wylaczenie;
    {przesun: -1|1} — w gore / w dol listy (zamiana z sasiadem)."""
    lic = db.get_or_404(LicencjaZwrotow, lic_id)
    d = json_body()
    if 'nazwa' in d:
        lic.nazwa = _nazwa_licencji(d.get('nazwa'))
    if 'aktywna' in d:
        if not isinstance(d['aktywna'], bool):
            abort(400, 'Pole „aktywna” musi być true/false.')
        lic.aktywna = d['aktywna']
    if 'przesun' in d:
        if d['przesun'] not in (-1, 1) or isinstance(d['przesun'], bool):
            abort(400, 'Pole „przesun” musi być -1 albo 1.')
        lista = LicencjaZwrotow.query.filter_by(rodzaj=lic.rodzaj)\
            .order_by(LicencjaZwrotow.kolejnosc, LicencjaZwrotow.id).all()
        i = lista.index(lic)
        j = i + d['przesun']
        if 0 <= j < len(lista):
            lista[i], lista[j] = lista[j], lista[i]
            for poz, element in enumerate(lista):
                element.kolejnosc = poz
    _zapisz_licencje()
    return jsonify(lic.to_dict())


@bp.route('/api/zwroty/licencje/<int:lic_id>', methods=['DELETE'])
@admin_required
def api_licencja_usun(lic_id):
    lic = db.get_or_404(LicencjaZwrotow, lic_id)
    # Przyciski Przyjecia zwrotu mozna usuwac zawsze — zwrot trzyma nazwe jako tekst.
    if SortingDataWpis.query.filter_by(licencja_id=lic.id).first():
        abort(409, 'Pozycja ma wpisane dane w Sorting Data — można ją wyłączyć albo zmienić nazwę, nie usunąć.')
    db.session.delete(lic)
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/admin/zwroty/cennik')
@admin_required
def admin_cennik():
    return render_template('admin_cennik_zwrotow.html', czynnosci=CZYNNOSCI_FAKTURY,
                           dzis=local_today())


@bp.route('/api/zwroty/cennik/<int:rok>/<int:miesiac>', methods=['GET'])
@admin_required
def api_cennik(rok, miesiac):
    _rok_miesiac(rok, miesiac)
    return jsonify(ceny_miesiaca(rok, miesiac))


@bp.route('/api/zwroty/cennik/<int:rok>/<int:miesiac>', methods=['PUT'])
@admin_required
def api_cennik_zapisz(rok, miesiac):
    """{ceny: {czynnosc: cena | null}} — null = brak ceny (Total charge pusty)."""
    _rok_miesiac(rok, miesiac)
    ceny = json_body().get('ceny')
    if not isinstance(ceny, dict):
        abort(400, 'Pole „ceny” musi być obiektem {czynność: cena}.')
    for czynnosc, cena in ceny.items():
        if czynnosc not in KLUCZE_FAKTURY:
            abort(400, f'Nieznana czynność: {czynnosc}.')
        if cena is not None and (isinstance(cena, bool) or not isinstance(cena, (int, float)) or cena < 0):
            abort(400, f'Cena dla {czynnosc} musi być liczbą ≥ 0.')
    zapisz_ceny(rok, miesiac, {k: (None if v is None else float(v)) for k, v in ceny.items()},
                current_user.id)
    return jsonify(ceny_miesiaca(rok, miesiac))
