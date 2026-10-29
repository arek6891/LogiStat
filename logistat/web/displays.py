"""Dekoracje (kafelek od 2026-10-24, do 2026-10-28 „Akcje - Displays”).

Akcje na wykresie Gantta (`/displays`), ekran akcji z indexami (autozapis),
historia zakonczonych akcji i listy typow / modeli stojakow (Supervisor).
Akcja zakonczona jest tylko do odczytu — zmienia ja dopiero „Przywróć”.
Etykiety POS: lista numerow z Excela → arkusz kodow kreskowych (PDF).
Stock stojakow: skan nosnik → strefa na stanowisku (Operacje magazynowe),
tabela stocku z POS / akcja / strefa i historia ruchow (lider).
"""
import io
import logging

import openpyxl
from flask import Blueprint, abort, flash, jsonify, make_response, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError

from logistat.czas import local_today, teraz_utc
from logistat.dekoracje import (
    DNI_OSTRZEZENIA_ETAPU, ETAPY_AKCJI, ETYKIET_MAX, KOPIE_MAX, LICZBA_MAX, NAGLOWEK_EKSPORTU_POS, NAZWA_MAX, NAZWA_SLOWNIKA_MAX,
    POLA_LICZB, POLA_PRODUKTYWNOSCI, POLA_SLOWNIKA, RODZAJE_SLOWNIKA, STREFA, data_z_tekstu,
    kod_nosnika, kod_strefy, wiersze_pos)
from logistat.dekoracje_dane import (
    akcja_dict, akcje_gantt, akcje_historia, etapy_akcji, importuj_pos, index_dict, listy_dla_akcji,
    ostatni_import_pos, ostatnie_ruchy, pos_dict, pos_lista, pos_w_stocku, przenies_stojak,
    ruch_dict, slownik, stock_stojakow, stojak_dict, strefa_po_kodzie, uzyte_pozycje_slownika,
    zapisz_etap)
from logistat.extensions import db
from logistat.models import (
    AkcjaDekoracji, EtapAkcji, IndexDekoracji, PosDekoracji, RuchStojaka, SlownikDekoracji, StojakDekoracji)
from logistat.pdf_etykiet import NA_STRONIE, pdf_etykiet
from logistat.przyjecie_zwrotu import pos_nummer
from logistat.web.excel import odpowiedz_xlsx, styl_naglowka
from logistat.web.wspolne import admin_required, json_body, leader_required, magazyn_required

bp = Blueprint('displays', __name__)
logger = logging.getLogger(__name__)


# ── ekrany ───────────────────────────────────────────────────────────────────

@bp.route('/displays')
@leader_required
def displays():
    dzis = local_today()
    return render_template('displays.html', akcje=akcje_gantt(dzis), dzis=dzis)


@bp.route('/displays/historia')
@leader_required
def historia():
    return render_template('displays_historia.html', akcje=akcje_historia(local_today()))


@bp.route('/displays/akcje/<int:akcja_id>')
@leader_required
def akcja(akcja_id):
    a = db.get_or_404(AkcjaDekoracji, akcja_id)
    return render_template('displays_akcja.html', akcja=akcja_dict(a, local_today()),
                           listy=listy_dla_akcji(a), pola_liczb=POLA_LICZB,
                           pola_prod=POLA_PRODUKTYWNOSCI, max_nazwy=NAZWA_MAX, max_liczby=LICZBA_MAX,
                           dni_ostrzezenia=DNI_OSTRZEZENIA_ETAPU)


# ── walidacja ────────────────────────────────────────────────────────────────

def _nazwa(wartosc, maks=NAZWA_MAX):
    nazwa = ' '.join(wartosc.split()) if isinstance(wartosc, str) else ''
    if not nazwa:
        abort(400, 'Nazwa nie może być pusta.')
    if len(nazwa) > maks:
        abort(400, f'Nazwa może mieć najwyżej {maks} znaków.')
    return nazwa


def _daty(d):
    """Obie daty naraz (ekran wysyla je razem) — inaczej zmiana „od” przed „do”
    zapisywalaby posredni, bledny stan."""
    od, do = data_z_tekstu(d.get('od')), data_z_tekstu(d.get('do'))
    if od is None or do is None:
        abort(400, 'Podaj datę startu i końca akcji (RRRR-MM-DD).')
    if do < od:
        abort(400, 'Data końca nie może być przed datą startu.')
    return od, do


def _liczba(pole, wartosc, calkowita):
    if wartosc is None or wartosc == '':
        return None
    if isinstance(wartosc, bool) or not isinstance(wartosc, (int, float)):
        abort(400, f'{pole}: wpisz liczbę.')
    if calkowita and wartosc != int(wartosc):
        abort(400, f'{pole}: wpisz liczbę całkowitą.')
    if wartosc < 0 or wartosc >= LICZBA_MAX:
        abort(400, f'{pole}: liczba od 0 do {LICZBA_MAX - 1}.')
    return int(wartosc) if calkowita else float(wartosc)


def _edytowalna(akcja):
    if akcja.zakonczona_at is not None:
        abort(409, 'Akcja jest zakończona — najpierw „Przywróć”, żeby ją zmienić.')


def _zmieniono(obiekt):
    obiekt.zmienil_id = current_user.id
    obiekt.zmieniono_at = teraz_utc()


# ── akcje ────────────────────────────────────────────────────────────────────

@bp.route('/api/dekoracje/akcje', methods=['POST'])
@leader_required
def api_akcja_dodaj():
    d = json_body()
    od, do = _daty(d)
    a = AkcjaDekoracji(nazwa=_nazwa(d.get('nazwa')), data_od=od, data_do=do,
                       utworzyl_id=current_user.id)
    db.session.add(a)
    db.session.commit()
    return jsonify(akcja_dict(a, local_today())), 201


@bp.route('/api/dekoracje/akcje/<int:akcja_id>', methods=['PUT'])
@leader_required
def api_akcja_zmien(akcja_id):
    """{nazwa, od, do} — nazwa i daty razem (jeden autozapis naglowka)."""
    a = db.get_or_404(AkcjaDekoracji, akcja_id)
    _edytowalna(a)
    d = json_body()
    a.nazwa = _nazwa(d.get('nazwa'))
    a.data_od, a.data_do = _daty(d)
    # Kamienie milowe musza zostac w datach akcji — najpierw przesuwa sie etap.
    nazwy = {k: n for k, n, *_ in ETAPY_AKCJI}
    poza = [f'{nazwy.get(e.etap, e.etap)} ({e.data_od:%d.%m}–{e.data_do:%d.%m})'
            for e in EtapAkcji.query.filter_by(akcja_id=a.id)
            if e.data_od and e.data_do and (e.data_od < a.data_od or e.data_do > a.data_do)]
    if poza:
        db.session.rollback()
        abort(400, 'Kamienie milowe wychodzą poza nowe daty akcji: ' + ', '.join(poza)
              + '. Najpierw przesuń ich terminy (przesuwając akcję później: wydłuż koniec, przesuń etapy, potem start).')
    _zmieniono(a)
    db.session.commit()
    return jsonify(akcja_dict(a, local_today(), z_indexami=False))


@bp.route('/api/dekoracje/akcje/<int:akcja_id>/status', methods=['PUT'])
@leader_required
def api_akcja_status(akcja_id):
    """{zakonczona: true} — „Zakończ akcję” (do historii); false — „Przywróć”."""
    a = db.get_or_404(AkcjaDekoracji, akcja_id)
    zakonczona = json_body().get('zakonczona')
    if not isinstance(zakonczona, bool):
        abort(400, 'Pole „zakonczona” musi być true/false.')
    if zakonczona and a.zakonczona_at is None:
        a.zakonczona_at, a.zakonczyl_id = teraz_utc(), current_user.id
    elif not zakonczona:
        a.zakonczona_at = a.zakonczyl_id = None
    _zmieniono(a)
    db.session.commit()
    return jsonify(akcja_dict(a, local_today(), z_indexami=False))


# ── kamienie milowe ──────────────────────────────────────────────────────────

@bp.route('/api/dekoracje/akcje/<int:akcja_id>/etapy/<etap>', methods=['PUT'])
@leader_required
def api_etap_zmien(akcja_id, etap):
    """{od, do} — obie daty razem (albo obie puste = bez terminu), wewnatrz dat
    akcji; {nie_dotyczy: bool} — tylko etap opcjonalny (Picking)."""
    a = db.get_or_404(AkcjaDekoracji, akcja_id)
    _edytowalna(a)
    opis = next((e for e in ETAPY_AKCJI if e[0] == etap), None)
    if opis is None:
        abort(404, 'Nieznany kamień milowy.')
    _, nazwa, _, opcjonalny = opis
    d = json_body()
    obecny = EtapAkcji.query.filter_by(akcja_id=a.id, etap=etap).first()
    od, do = (obecny.data_od, obecny.data_do) if obecny else (None, None)
    nie_dotyczy = bool(obecny and obecny.nie_dotyczy)
    if 'od' in d or 'do' in d:
        puste = [d.get(k) in (None, '') for k in ('od', 'do')]
        if all(puste):
            od = do = None
        else:
            od, do = data_z_tekstu(d.get('od')), data_z_tekstu(d.get('do'))
            if od is None or do is None:
                abort(400, f'{nazwa}: podaj obie daty (od i do) albo zostaw obie puste.')
            if do < od:
                abort(400, f'{nazwa}: data końca nie może być przed datą startu.')
            if od < a.data_od or do > a.data_do:
                abort(400, f'{nazwa}: termin musi mieścić się w datach akcji '
                           f'({a.data_od:%d.%m.%Y} – {a.data_do:%d.%m.%Y}).')
    if 'nie_dotyczy' in d:
        if not isinstance(d['nie_dotyczy'], bool):
            abort(400, 'Pole „nie_dotyczy” musi być true/false.')
        if d['nie_dotyczy'] and not opcjonalny:
            abort(400, f'{nazwa}: ten etap jest zawsze potrzebny.')
        nie_dotyczy = d['nie_dotyczy']
    if nie_dotyczy:
        # Etap pominiety nie ma terminu — inaczej stare daty blokowalyby
        # zmiane dat akcji, a pola dat sa wtedy wylaczone.
        od = do = None
    zapisz_etap(a.id, etap, od, do, nie_dotyczy, current_user.id)
    db.session.commit()
    return jsonify(next(e for e in etapy_akcji([a.id], local_today())[a.id] if e['klucz'] == etap))


# ── indexy ───────────────────────────────────────────────────────────────────

def _pozycja_slownika(pole, rodzaj, wartosc, obecna):
    if wartosc is None or wartosc == '':
        return None
    if isinstance(wartosc, bool) or not isinstance(wartosc, int):
        abort(400, f'Nieznana pozycja listy „{RODZAJE_SLOWNIKA[rodzaj]}”.')
    p = db.session.get(SlownikDekoracji, wartosc)
    # Wylaczona pozycja zostaje mozliwa tylko w indexie, ktory juz ja ma.
    if p is None or p.rodzaj != rodzaj or (not p.aktywna and p.id != obecna):
        abort(400, f'Nieznana albo wyłączona pozycja listy „{RODZAJE_SLOWNIKA[rodzaj]}” — odśwież stronę.')
    return p.id


@bp.route('/api/dekoracje/akcje/<int:akcja_id>/indexy', methods=['POST'])
@leader_required
def api_index_dodaj(akcja_id):
    """„+ Index” — pusty index z nazwa „Index N” (nazwe i reszte wpisuje sie potem)."""
    a = db.get_or_404(AkcjaDekoracji, akcja_id)
    _edytowalna(a)
    nr = (db.session.query(db.func.max(IndexDekoracji.kolejnosc))
          .filter(IndexDekoracji.akcja_id == a.id).scalar() or 0) + 1
    ix = IndexDekoracji(akcja_id=a.id, nazwa=f'Index {nr}', kolejnosc=nr)
    _zmieniono(ix)
    db.session.add(ix)
    db.session.commit()
    return jsonify(index_dict(ix)), 201


@bp.route('/api/dekoracje/indexy/<int:index_id>', methods=['PUT'])
@leader_required
def api_index_zmien(index_id):
    """Czesciowy zapis — zmienia tylko podane pola (dwie osoby na roznych
    polach jednego indexu nie nadpisuja sobie nawzajem)."""
    ix = db.get_or_404(IndexDekoracji, index_id)
    _edytowalna(ix.akcja)
    d = json_body()
    etykiety = dict(POLA_LICZB + POLA_PRODUKTYWNOSCI)
    if 'nazwa' in d:
        ix.nazwa = _nazwa(d['nazwa'])
    for pole, _ in POLA_LICZB:
        if pole in d:
            setattr(ix, pole, _liczba(etykiety[pole], d[pole], calkowita=True))
    for pole, _ in POLA_PRODUKTYWNOSCI:
        if pole in d:
            setattr(ix, pole, _liczba(etykiety[pole], d[pole], calkowita=False))
    for pole, rodzaj in POLA_SLOWNIKA:
        if pole in d:
            setattr(ix, pole, _pozycja_slownika(pole, rodzaj, d[pole], getattr(ix, pole)))
    _zmieniono(ix)
    db.session.commit()
    return jsonify(index_dict(ix))


@bp.route('/api/dekoracje/indexy/<int:index_id>', methods=['DELETE'])
@leader_required
def api_index_usun(index_id):
    ix = db.get_or_404(IndexDekoracji, index_id)
    _edytowalna(ix.akcja)
    db.session.delete(ix)
    db.session.commit()
    return jsonify({'ok': True})


# ── listy (Supervisor → Listy Dekoracji) ─────────────────────────────────────

@bp.route('/admin/dekoracje')
@admin_required
def admin_slownik():
    lista = slownik()
    return render_template('admin_dekoracje.html', rodzaje=RODZAJE_SLOWNIKA,
                           licencje={r: [p for p in lista if p.rodzaj == r] for r in RODZAJE_SLOWNIKA},
                           uzywane=uzyte_pozycje_slownika(), max_nazwy=NAZWA_SLOWNIKA_MAX)


def _nazwa_slownika(rodzaj, wartosc):
    nazwa = _nazwa(wartosc, NAZWA_SLOWNIKA_MAX)
    if rodzaj != STREFA:
        return nazwa
    kod = kod_strefy(nazwa)
    if kod is None:
        abort(400, 'Nazwa strefy to kod na jej etykiecie: litery A–Z, cyfry i - _ / . '
                   '(bez spacji i polskich liter), najwyżej kilkanaście znaków.')
    return kod


def _zapisz_slownik():
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        abort(409, 'Taka nazwa już jest na tej liście.')


@bp.route('/api/dekoracje/slownik', methods=['POST'])
@admin_required
def api_slownik_dodaj():
    d = json_body()
    rodzaj = d.get('rodzaj')
    if rodzaj not in RODZAJE_SLOWNIKA:
        abort(400, 'Nieznana lista.')
    nazwa = _nazwa_slownika(rodzaj, d.get('nazwa'))
    ostatnia = db.session.query(db.func.max(SlownikDekoracji.kolejnosc))\
        .filter(SlownikDekoracji.rodzaj == rodzaj).scalar()
    p = SlownikDekoracji(rodzaj=rodzaj, nazwa=nazwa, kolejnosc=(ostatnia or 0) + 1)
    db.session.add(p)
    _zapisz_slownik()
    return jsonify(p.to_dict()), 201


@bp.route('/api/dekoracje/slownik/<int:poz_id>', methods=['PUT'])
@admin_required
def api_slownik_zmien(poz_id):
    """{nazwa} / {aktywna} / {przesun: -1|1} — jak listy Sorting Data."""
    p = db.get_or_404(SlownikDekoracji, poz_id)
    d = json_body()
    if 'nazwa' in d:
        p.nazwa = _nazwa_slownika(p.rodzaj, d.get('nazwa'))
    if 'aktywna' in d:
        if not isinstance(d['aktywna'], bool):
            abort(400, 'Pole „aktywna” musi być true/false.')
        p.aktywna = d['aktywna']
    if 'przesun' in d:
        if d['przesun'] not in (-1, 1) or isinstance(d['przesun'], bool):
            abort(400, 'Pole „przesun” musi być -1 albo 1.')
        lista = slownik(p.rodzaj)
        i = lista.index(p)
        j = i + d['przesun']
        if 0 <= j < len(lista):
            lista[i], lista[j] = lista[j], lista[i]
            for poz, element in enumerate(lista):
                element.kolejnosc = poz
    _zapisz_slownik()
    return jsonify(p.to_dict())


@bp.route('/api/dekoracje/slownik/<int:poz_id>', methods=['DELETE'])
@admin_required
def api_slownik_usun(poz_id):
    p = db.get_or_404(SlownikDekoracji, poz_id)
    if p.id in uzyte_pozycje_slownika():
        abort(409, 'Pozycja jest użyta w indexie akcji — można ją wyłączyć albo zmienić nazwę, nie usunąć.')
    db.session.delete(p)
    db.session.commit()
    return jsonify({'ok': True})


# ── Etykiety POS ─────────────────────────────────────────────────────────────

@bp.route('/displays/pos')
@leader_required
def pos():
    lista = pos_lista()
    return render_template('displays_pos.html', lista=[pos_dict(p) for p in lista],
                           ostatni_import=len(ostatni_import_pos()), na_stronie=NA_STRONIE,
                           kopie_max=KOPIE_MAX, naglowki=NAGLOWEK_EKSPORTU_POS)


@bp.route('/api/dekoracje/pos/import', methods=['POST'])
@leader_required
def api_pos_import():
    """Plik .xlsx: kolumna POS (wymagana) i Post Processing (opcjonalna)."""
    plik = request.files.get('file')
    if plik is None or not plik.filename:
        abort(400, 'Wybierz plik Excel (.xlsx).')
    try:
        wb = openpyxl.load_workbook(io.BytesIO(plik.read()), data_only=True, read_only=True)
    except Exception:
        abort(400, 'Nie można odczytać pliku — potrzebny Excel .xlsx.')
    try:
        ws = wb.active
        if ws is None:
            abort(400, 'Plik Excel jest pusty.')
        pozycje, ma_pp, bledy = wiersze_pos(ws.iter_rows(values_only=True))
    except ValueError as e:
        abort(400, str(e))
    finally:
        wb.close()
    if not pozycje and not bledy:
        abort(400, 'W pliku nie ma numerów POS.')
    wynik = importuj_pos(pozycje, ma_pp, current_user.id)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        abort(409, 'Ktoś w tej chwili importował te same numery — wczytaj plik jeszcze raz.')
    logger.info('Import listy POS: %s, bledow %d (uzytkownik %s)', wynik, len(bledy), current_user.id)
    return jsonify({**wynik, 'wczytane': len(pozycje), 'bledy': bledy, 'ma_pp': ma_pp,
                    'lista': [pos_dict(p) for p in pos_lista()],
                    'ostatni_import': len(ostatni_import_pos())})


@bp.route('/api/dekoracje/pos/<int:pos_id>', methods=['DELETE'])
@leader_required
def api_pos_usun(pos_id):
    p = db.get_or_404(PosDekoracji, pos_id)
    if p.id in pos_w_stocku():
        abort(409, f'POS {p.numer} jest przypisany do stojaka w stocku — najpierw zmień POS stojaka.')
    db.session.delete(p)
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/displays/pos/etykiety', methods=['POST'])
@leader_required
def pos_etykiety():
    """Arkusz etykiet (PDF w nowej karcie). Formularz: `zakres` = wszystkie |
    import (ostatni) | wybrane (+ `id`), `kopie` kazdej etykiety."""
    zakres = request.form.get('zakres')
    if zakres == 'wszystkie':
        lista = pos_lista()
    elif zakres == 'import':
        lista = ostatni_import_pos()
    elif zakres == 'wybrane':
        ids = {int(i) for i in request.form.getlist('id') if i.isdigit()}
        lista = [p for p in pos_lista() if p.id in ids]
    else:
        lista = []
    try:
        kopie = int(request.form.get('kopie') or 1)
    except ValueError:
        kopie = 0
    if not 1 <= kopie <= KOPIE_MAX:
        flash(f'Liczba kopii od 1 do {KOPIE_MAX}.', 'error')
        return redirect(url_for('displays.pos'))
    if not lista:
        flash('Nie ma etykiet do druku — zaznacz numery POS albo wczytaj listę.', 'error')
        return redirect(url_for('displays.pos'))
    if len(lista) * kopie > ETYKIET_MAX:
        flash(f'Za dużo etykiet naraz ({len(lista) * kopie}) — najwyżej {ETYKIET_MAX}.', 'error')
        return redirect(url_for('displays.pos'))
    odpowiedz = make_response(pdf_etykiet([p.numer for p in lista for _ in range(kopie)]))
    odpowiedz.headers['Content-Type'] = 'application/pdf'
    odpowiedz.headers['Content-Disposition'] = f'inline; filename="etykiety-pos-{local_today():%Y-%m-%d}.pdf"'
    return odpowiedz


@bp.route('/displays/pos/export')
@leader_required
def pos_export():
    """Lista do poprawki w Excelu — ten sam uklad wczytuje import (klucz = POS)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = 'POS'
    ws.append(list(NAGLOWEK_EKSPORTU_POS))
    for komorka in ws[1]:
        styl_naglowka(komorka, '325C96')
    for p in pos_lista():
        ws.append([p.numer, {True: 'Tak', False: 'Nie'}.get(p.post_processing)])
        # Numer zawsze jako tekst — inaczej Excel zjada wiodace zera.
        ws.cell(ws.max_row, 1).number_format = '@'
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 18
    ws.auto_filter.ref = ws.dimensions
    return odpowiedz_xlsx(wb, f'POS_{local_today():%Y-%m-%d}.xlsx')


# ── Stock stojakow: stanowisko (Operacje magazynowe) ─────────────────────────

def _nosnik(wartosc):
    kod = kod_nosnika(wartosc)
    if kod is None:
        abort(400, 'Nieprawidłowy numer nośnika — zeskanuj etykietę nośnika.')
    if strefa_po_kodzie(kod, tylko_aktywne=False) is not None:
        abort(400, f'„{kod}” to kod strefy — najpierw zeskanuj nośnik.')
    return kod


def _stojak_z_ruchem(s):
    return stojak_dict(s, ostatnie_ruchy([s.id]).get(s.id))


@bp.route('/dekoracje/skan')
@magazyn_required
def skan_stojakow():
    strefy = [p.nazwa for p in slownik(STREFA) if p.aktywna]
    return render_template('dekoracje_skan.html', strefy=strefy)


@bp.route('/api/dekoracje/nosnik')
@magazyn_required
def api_nosnik():
    """Krok 1 stanowiska: co wiadomo o nosniku (nowy = zalozy go zapis)."""
    kod = _nosnik(request.args.get('kod'))
    s = StojakDekoracji.query.filter_by(nosnik=kod).first()
    return jsonify({'nowy': s is None, 'nosnik': kod, 'stojak': _stojak_z_ruchem(s) if s else None})


@bp.route('/api/dekoracje/skan', methods=['POST'])
@magazyn_required
def api_skan():
    """{nosnik, strefa} — nosnik odlozony w strefie. Nieznany nosnik zaklada
    stojak; zmiana strefy zapisuje ruch (ta sama strefa → bez ruchu)."""
    d = json_body()
    kod = _nosnik(d.get('nosnik'))
    kod_s = kod_strefy(d.get('strefa')) if isinstance(d.get('strefa'), str) else None
    strefa = strefa_po_kodzie(kod_s) if kod_s else None
    if strefa is None:
        abort(400, 'Nieznana strefa — zeskanuj etykietę strefy (lista: Supervisor → Listy Dekoracji).')
    teraz = teraz_utc()
    nowy = False
    s = StojakDekoracji.query.filter_by(nosnik=kod).with_for_update().first()
    if s is None:
        s = StojakDekoracji(nosnik=kod, utworzyl_id=current_user.id, utworzono_at=teraz)
        db.session.add(s)
        try:
            db.session.flush()
            nowy = True
        except IntegrityError:
            # Drugie stanowisko zalozylo ten nosnik w tej samej chwili.
            db.session.rollback()
            s = StojakDekoracji.query.filter_by(nosnik=kod).with_for_update().one()
    przeniesiono = przenies_stojak(s, strefa.id, current_user.id, recznie=False, teraz=teraz)
    db.session.commit()
    return jsonify({'nowy': nowy, 'przeniesiono': przeniesiono, 'stojak': _stojak_z_ruchem(s)})


# ── Stock stojakow: tabela lidera ────────────────────────────────────────────

@bp.route('/displays/stock')
@leader_required
def stock():
    stojaki = stock_stojakow()
    uzyte_akcje = {s['akcja_id'] for s in stojaki if s['akcja_id']}
    akcje = AkcjaDekoracji.query.filter(db.or_(AkcjaDekoracji.zakonczona_at.is_(None),
                                               AkcjaDekoracji.id.in_(uzyte_akcje or [0])))\
        .order_by(AkcjaDekoracji.data_od, AkcjaDekoracji.id).all()
    uzyte_strefy = {s['strefa_id'] for s in stojaki if s['strefa_id']}
    return render_template('displays_stock.html', stojaki=stojaki,
                           akcje=[{'id': a.id, 'nazwa': a.nazwa, 'zakonczona': a.zakonczona_at is not None}
                                  for a in akcje],
                           strefy=[{'id': p.id, 'nazwa': p.nazwa, 'aktywna': p.aktywna}
                                   for p in slownik(STREFA) if p.aktywna or p.id in uzyte_strefy],
                           pos=[p.numer for p in pos_lista()])


@bp.route('/api/dekoracje/stojaki/<int:stojak_id>', methods=['PUT'])
@leader_required
def api_stojak_zmien(stojak_id):
    """Czesciowy zapis: {pos: numer | ''} (tylko z listy Etykiet POS),
    {akcja_id: id | null} (niezakonczona albo obecna), {strefa_id} (ruch „recznie”)."""
    s = db.get_or_404(StojakDekoracji, stojak_id)
    d = json_body()
    if 'pos' in d:
        if d['pos'] in (None, ''):
            s.pos_id = None
        else:
            numer = pos_nummer(d['pos']) if isinstance(d['pos'], str) else None
            p = PosDekoracji.query.filter_by(numer=numer).first() if numer else None
            if p is None:
                abort(400, 'Tego numeru nie ma na liście Etykiet POS — najpierw wczytaj go do listy.')
            s.pos_id = p.id
    if 'akcja_id' in d:
        akcja_id = d['akcja_id']
        if akcja_id is None:
            s.akcja_id = None
        else:
            a = db.session.get(AkcjaDekoracji, akcja_id) \
                if isinstance(akcja_id, int) and not isinstance(akcja_id, bool) else None
            if a is None or (a.zakonczona_at is not None and a.id != s.akcja_id):
                abort(400, 'Nieznana albo zakończona akcja — odśwież stronę.')
            s.akcja_id = a.id
    if 'strefa_id' in d:
        strefa_id = d['strefa_id']
        p = db.session.get(SlownikDekoracji, strefa_id) \
            if isinstance(strefa_id, int) and not isinstance(strefa_id, bool) else None
        if p is None or p.rodzaj != STREFA or (not p.aktywna and p.id != s.strefa_id):
            abort(400, 'Nieznana albo wyłączona strefa — odśwież stronę.')
        przenies_stojak(s, p.id, current_user.id, recznie=True)
    _zmieniono(s)
    db.session.commit()
    return jsonify(_stojak_z_ruchem(s))


@bp.route('/api/dekoracje/stojaki/<int:stojak_id>', methods=['DELETE'])
@leader_required
def api_stojak_usun(stojak_id):
    """Pomylkowy skan — stojak znika razem z historia ruchow."""
    s = db.get_or_404(StojakDekoracji, stojak_id)
    logger.info('Usuniety stojak %s (nosnik %s) przez %s', s.id, s.nosnik, current_user.id)
    db.session.delete(s)
    db.session.commit()
    return jsonify({'ok': True})


@bp.route('/api/dekoracje/stojaki/<int:stojak_id>/ruchy')
@leader_required
def api_stojak_ruchy(stojak_id):
    s = db.get_or_404(StojakDekoracji, stojak_id)
    ruchy = RuchStojaka.query.filter_by(stojak_id=s.id)\
        .order_by(RuchStojaka.kiedy.desc(), RuchStojaka.id.desc()).all()
    return jsonify([ruch_dict(r) for r in ruchy])


@bp.route('/admin/dekoracje/strefy/etykiety', methods=['POST'])
@admin_required
def strefy_etykiety():
    """Etykiety aktywnych stref (do powieszenia w hali) — ten sam arkusz co POS."""
    strefy = [p.nazwa for p in slownik(STREFA) if p.aktywna]
    if not strefy:
        flash('Brak aktywnych stref — dodaj je w liście „Strefy”.', 'error')
        return redirect(url_for('displays.admin_slownik'))
    odpowiedz = make_response(pdf_etykiet(strefy))
    odpowiedz.headers['Content-Type'] = 'application/pdf'
    odpowiedz.headers['Content-Disposition'] = 'inline; filename="etykiety-stref.pdf"'
    return odpowiedz
