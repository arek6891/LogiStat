"""Invoice specification i Sorting Data (kafelek Zwroty, 2026-10-26).

- uklad i nazwy 1:1 z zakladek arkusza operacji, eksport do Excela z formulami;
- ilosci wpisuje lider (puste ≠ 0), Sorting returns liczy system z Zwroty Paczek
  (liczba WIERSZY raportu w miesiacu), ceny ustawia admin miesiacami;
- lider nie widzi cen ani Total charge — ani na ekranie, ani w zrodle, ani w Excelu;
- listy wierszy Sorting Data: admin, pozycja z danymi nie znika z historii.
Wszystkie liczby syntetyczne.
"""
import io
from datetime import date, datetime

import openpyxl
import pytest

from logistat.extensions import db
from logistat.faktura_zwrotow import KLUCZ_POS, LICENCJA, LICENCJE_DOMYSLNE, WYSYLKA
from logistat.models import (
    CenaZwrotow, IloscMiesiacaZwrotow, LicencjaZwrotow, SortingDataWpis, ZwrotPaczki)
from logistat.zwroty_dane import invoice_specification, sorting_data

CENA_TESTOWA = 7.37  # nietypowa, zeby jej szukac w HTML / Excelu lidera


def ustaw_cene(czynnosc='palletizing', cena=CENA_TESTOWA, rok=2026, miesiac=9):
    db.session.add(CenaZwrotow(rok=rok, miesiac=miesiac, czynnosc=czynnosc, cena=cena))
    db.session.commit()


def xlsx(odpowiedz):
    assert odpowiedz.status_code == 200, odpowiedz.get_data(as_text=True)[:300]
    return openpyxl.load_workbook(io.BytesIO(odpowiedz.data)).active


# ── Invoice specification ────────────────────────────────────────────────────

def test_ilosc_wpisana_przez_lidera_i_puste_to_nie_zero(leader_client):
    r = leader_client.put('/api/zwroty/invoice-specification/ilosc',
                          json={'rok': 2026, 'miesiac': 9, 'klucz': 'palletizing', 'ilosc': 12})
    assert r.status_code == 200, r.get_json()
    r = leader_client.put('/api/zwroty/invoice-specification/ilosc',
                          json={'rok': 2026, 'miesiac': 9, 'klucz': 'loading', 'ilosc': 0})
    assert r.status_code == 200

    w = {x['klucz']: x for x in invoice_specification(2026, 9)['wiersze']}
    assert w['palletizing']['ilosc'] == 12
    assert w['loading']['ilosc'] == 0
    assert w['unloading']['ilosc'] is None

    leader_client.put('/api/zwroty/invoice-specification/ilosc',
                      json={'rok': 2026, 'miesiac': 9, 'klucz': 'palletizing', 'ilosc': None})
    w = {x['klucz']: x for x in invoice_specification(2026, 9)['wiersze']}
    assert w['palletizing']['ilosc'] is None, 'wyczyszczone pole = puste, nie 0'
    assert IloscMiesiacaZwrotow.query.count() == 2, 'upsert, nie nowe wiersze'


@pytest.mark.parametrize('body,blad', [
    ({'rok': 2026, 'miesiac': 13, 'klucz': 'loading', 'ilosc': 1}, 'miesiac'),
    ({'rok': 2026, 'miesiac': 9, 'klucz': 'nieznana', 'ilosc': 1}, 'czynność'),
    ({'rok': 2026, 'miesiac': 9, 'klucz': 'loading', 'ilosc': -1}, 'od 0'),
    ({'rok': 2026, 'miesiac': 9, 'klucz': 'loading', 'ilosc': 1.5}, 'całkowitą'),
    ({'rok': 2026, 'miesiac': 9, 'klucz': 'loading', 'ilosc': True}, 'całkowitą'),
    ({'rok': 2026, 'miesiac': 9, 'klucz': 'sorting_returns', 'ilosc': 5}, 'Zwroty Paczek'),
])
def test_walidacja_ilosci(leader_client, body, blad):
    r = leader_client.put('/api/zwroty/invoice-specification/ilosc', json=body)
    assert r.status_code == 400 and blad in r.get_json()['error']


def test_sorting_returns_to_liczba_wierszy_raportu_zwrotow(flask_app):
    """Arkusz: 1816 wierszy (4 Paketnummer powtorzone) → Sorting returns 1816.
    Granice miesiaca po czasie lokalnym (Warszawa), nie UTC."""
    def zwrot(token, kiedy_utc, paket='%A'):
        db.session.add(ZwrotPaczki(
            paketnummer=paket, brand='ACC', pos_nummer='NO', kurier='DPD', retoure=1,
            verpackungsvorschrift='Yes', token=token, przyjeto_at=kiedy_utc))

    zwrot('t-luty-1', datetime(2026, 2, 10, 9))
    zwrot('t-luty-2', datetime(2026, 2, 10, 10))          # ten sam Paketnummer — liczy sie znowu
    zwrot('t-luty-3', datetime(2026, 2, 28, 22, 30), '%B')  # 28.02 23:30 w Warszawie
    zwrot('t-marzec', datetime(2026, 2, 28, 23, 30))        # 01.03 00:30 w Warszawie
    zwrot('t-styczen', datetime(2026, 1, 31, 22, 59))       # 31.01 23:59 w Warszawie
    db.session.commit()

    w = invoice_specification(2026, 2)['wiersze'][0]
    assert (w['klucz'], w['ilosc'], w['auto']) == ('sorting_returns', 3, True)


def test_total_charge_i_brak_ceny(flask_app):
    db.session.add(IloscMiesiacaZwrotow(rok=2026, miesiac=9, klucz='palletizing', wartosc=10))
    db.session.commit()
    ustaw_cene('palletizing', 2.5)
    ustaw_cene('unloading', 4.0)

    d = invoice_specification(2026, 9)
    w = {x['klucz']: x for x in d['wiersze']}

    assert w['palletizing']['charge'] == 25.0
    assert w['unloading']['charge'] == 0.0, 'ilosc pusta, cena jest → 0'
    assert w['loading']['charge'] is None, 'bez ceny Total charge pusty, nie 0'
    assert d['total'] == 25.0
    assert 'Loading pallets' in d['brak_cen'] and 'Palletizing sorted articles' not in d['brak_cen']


def test_admin_widzi_ceny(admin_client):
    ustaw_cene()
    html = admin_client.get('/zwroty/invoice-specification?rok=2026&miesiac=9').get_data(as_text=True)
    assert 'Total charge' in html and '7.37' in html and '>Price<' in html
    assert 'Brak ceny' in html, 'czynnosci bez ceny sa wypisane, zeby nie wyszla zerowa faktura'


def test_lider_nie_widzi_cen(leader_client):
    ustaw_cene()
    html = leader_client.get('/zwroty/invoice-specification?rok=2026&miesiac=9').get_data(as_text=True)
    assert '7.37' not in html, 'cena nie moze trafic nawet do zrodla strony lidera'
    assert 'Total charge' not in html and '>Price<' not in html
    assert 'Warehousing Activities' in html and 'Returns logistics service' in html


def test_excel_admina_jak_arkusz_z_formulami(admin_client):
    ustaw_cene()
    ws = xlsx(admin_client.get('/zwroty/invoice-specification/export?rok=2026&miesiac=9'))

    assert ws['B3'].value == 'Warehousing Activities'
    assert ws['B4'].value == 'Returns logistics service'
    assert (ws['B6'].value, ws['C6'].value, ws['B7'].value, ws['C7'].value) == ('Year', 2026, 'Month', 9)
    assert [ws.cell(9, k).value for k in range(2, 7)] == \
        ['Activity', 'Total quantity', 'UoM', 'Price', 'Total charge']
    assert [ws.cell(r, 2).value for r in range(11, 16)] == [
        'Sorting returns', 'Palletizing sorted articles', 'Unloading pallets',
        'Taking weight of pallets', 'Loading pallets']
    assert [ws.cell(r, 4).value for r in range(11, 16)] == ['box', 'box', 'pallet', 'pallet', 'pallet']
    assert ws['E12'].value == CENA_TESTOWA
    assert ws['F11'].value == '=E11*C11'
    assert (ws['B17'].value, ws['F17'].value) == ('Total charge', '=SUM(F11:F15)')


def test_excel_lidera_bez_cen(leader_client):
    ustaw_cene()
    ws = xlsx(leader_client.get('/zwroty/invoice-specification/export?rok=2026&miesiac=9'))

    wartosci = [c.value for wiersz in ws.iter_rows() for c in wiersz if c.value is not None]
    assert CENA_TESTOWA not in wartosci
    assert 'Price' not in wartosci and 'Total charge' not in wartosci
    assert 'Palletizing sorted articles' in wartosci


# ── cennik (admin) ───────────────────────────────────────────────────────────

def test_cennik_zapis_odczyt_i_usuniecie_ceny(admin_client):
    r = admin_client.put('/api/zwroty/cennik/2026/9', json={'ceny': {'palletizing': 1.5, 'loading': 2}})
    assert r.status_code == 200 and r.get_json() == {'palletizing': 1.5, 'loading': 2.0}
    r = admin_client.put('/api/zwroty/cennik/2026/9', json={'ceny': {'palletizing': None, 'loading': 3}})
    assert r.get_json() == {'loading': 3.0}, 'null usuwa cene (brak ceny), a nie wpisuje 0'
    assert admin_client.get('/api/zwroty/cennik/2026/10').get_json() == {}


@pytest.mark.parametrize('ceny', [{'xxx': 1}, {'loading': -1}, {'loading': True}, {'loading': '2'}, []])
def test_cennik_walidacja(admin_client, ceny):
    assert admin_client.put('/api/zwroty/cennik/2026/9', json={'ceny': ceny}).status_code == 400


def test_cennik_tylko_admin(leader_client):
    assert leader_client.get('/api/zwroty/cennik/2026/9').status_code == 403
    assert leader_client.put('/api/zwroty/cennik/2026/9', json={'ceny': {'loading': 1}}).status_code == 403
    assert leader_client.get('/admin/zwroty/cennik').status_code == 302


# ── Sorting Data ─────────────────────────────────────────────────────────────

def licencja(nazwa, rodzaj=LICENCJA):
    return LicencjaZwrotow.query.filter_by(rodzaj=rodzaj, nazwa=nazwa).one()


def test_listy_zasiane_z_arkusza(flask_app):
    assert [lic.nazwa for lic in LicencjaZwrotow.query.filter_by(rodzaj=LICENCJA)
            .order_by(LicencjaZwrotow.kolejnosc)] == list(LICENCJE_DOMYSLNE[LICENCJA])
    assert licencja('RECYKLING').nazwa == 'RECYKLING', 'spacja z arkusza obcieta'
    assert LicencjaZwrotow.query.filter_by(rodzaj=WYSYLKA).count() == len(LICENCJE_DOMYSLNE[WYSYLKA])


def test_komorki_suma_i_pos(leader_client):
    disney, marvel = licencja('DISNEY'), licencja('Marvel')
    for lic, pole, w in ((disney, 'cartons', 48), (disney, 'pallets', 2), (marvel, 'cartons', 24)):
        r = leader_client.put('/api/zwroty/sorting-data', json={
            'rok': 2026, 'miesiac': 9, 'licencja_id': lic.id, 'pole': pole, 'wartosc': w})
        assert r.status_code == 200, r.get_json()
    wys = licencja('Disney boxes shipped', WYSYLKA)
    leader_client.put('/api/zwroty/sorting-data', json={
        'rok': 2026, 'miesiac': 9, 'licencja_id': wys.id, 'pole': 'cartons', 'wartosc': 100})
    leader_client.put('/api/zwroty/sorting-data', json={
        'rok': 2026, 'miesiac': 9, 'pole': KLUCZ_POS, 'wartosc': 321})

    d = sorting_data(2026)
    assert d['suma'][9] == {'cartons': 72, 'pallets': 2}, 'Total boxes built = tylko licencje'
    assert d['pos'][9] == 321 and d['pos'][8] is None
    wiersz = next(w for w in d['licencje'] if w['id'] == marvel.id)
    assert wiersz['mies'][9] == {'cartons': 24, 'pallets': None}


def test_sorting_walidacja(leader_client):
    lic = licencja('DISNEY')
    for body in ({'rok': 2026, 'miesiac': 9, 'licencja_id': lic.id, 'pole': 'xxx', 'wartosc': 1},
                 {'rok': 2026, 'miesiac': 9, 'licencja_id': 99999, 'pole': 'cartons', 'wartosc': 1},
                 {'rok': 2026, 'miesiac': 0, 'licencja_id': lic.id, 'pole': 'cartons', 'wartosc': 1},
                 {'rok': 2026, 'miesiac': 9, 'licencja_id': lic.id, 'pole': 'cartons', 'wartosc': 'abc'}):
        assert leader_client.put('/api/zwroty/sorting-data', json=body).status_code == 400, body


def test_wylaczona_licencja_z_danymi_zostaje_w_swoim_roku(flask_app):
    lic = licencja('Peppa')
    db.session.add(SortingDataWpis(licencja_id=lic.id, rok=2026, miesiac=1, cartons=5))
    lic.aktywna = False
    db.session.commit()

    assert lic.id in [w['id'] for w in sorting_data(2026)['licencje']]
    assert lic.id not in [w['id'] for w in sorting_data(2027)['licencje']]


def test_ekran_i_excel_sorting_data(leader_client):
    lic = licencja('DISNEY')
    db.session.add(SortingDataWpis(licencja_id=lic.id, rok=2026, miesiac=2, cartons=24, pallets=1))
    db.session.add(IloscMiesiacaZwrotow(rok=2026, miesiac=2, klucz=KLUCZ_POS, wartosc=55))
    db.session.commit()

    html = leader_client.get('/zwroty/sorting-data?rok=2026').get_data(as_text=True)
    for tekst in ('January', 'December', 'Cartons', 'Pallets', 'Total boxes built', 'POS processed',
                  'DISNEY', 'Disney boxes shipped'):
        assert tekst in html, tekst

    ws = xlsx(leader_client.get('/zwroty/sorting-data/export?rok=2026'))
    assert (ws['B2'].value, ws['D2'].value, ws['X2'].value) == ('January', 'February', 'December')
    assert (ws['B3'].value, ws['C3'].value) == ('Cartons', 'Pallets')
    kolumna_a = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    nr_disney = kolumna_a.index('DISNEY') + 1
    assert (ws.cell(nr_disney, 4).value, ws.cell(nr_disney, 5).value) == (24, 1)
    nr_sumy = kolumna_a.index('Total boxes built') + 1
    assert ws.cell(nr_sumy, 4).value == f'=SUM(D4:D{nr_sumy - 1})'
    nr_pos = kolumna_a.index('POS processed') + 1
    assert nr_pos == ws.max_row and ws.cell(nr_pos, 4).value == 55


# ── listy Sorting Data (admin) ───────────────────────────────────────────────

def test_admin_dodaje_przesuwa_wylacza_i_usuwa(admin_client):
    r = admin_client.post('/api/zwroty/licencje', json={'rodzaj': LICENCJA, 'nazwa': '  Nowa  '})
    assert r.status_code == 201 and r.get_json()['nazwa'] == 'Nowa'
    nowa_id = r.get_json()['id']
    assert admin_client.post('/api/zwroty/licencje', json={'rodzaj': LICENCJA, 'nazwa': 'Nowa'}).status_code == 409

    admin_client.put(f'/api/zwroty/licencje/{nowa_id}', json={'przesun': -1})
    nazwy = [lic.nazwa for lic in LicencjaZwrotow.query.filter_by(rodzaj=LICENCJA)
             .order_by(LicencjaZwrotow.kolejnosc)]
    assert nazwy[-2:] == ['Nowa', 'Other licenses']

    assert admin_client.put(f'/api/zwroty/licencje/{nowa_id}', json={'aktywna': False}).get_json()['aktywna'] is False
    assert admin_client.delete(f'/api/zwroty/licencje/{nowa_id}').status_code == 200


def test_pozycji_z_danymi_nie_da_sie_usunac(admin_client):
    lic = licencja('SILVER')
    db.session.add(SortingDataWpis(licencja_id=lic.id, rok=2026, miesiac=3, pallets=1))
    db.session.commit()

    assert admin_client.delete(f'/api/zwroty/licencje/{lic.id}').status_code == 409
    assert admin_client.put(f'/api/zwroty/licencje/{lic.id}', json={'nazwa': 'SILVER 925'}).status_code == 200


def test_listy_tylko_admin(leader_client):
    assert leader_client.post('/api/zwroty/licencje', json={'rodzaj': LICENCJA, 'nazwa': 'X'}).status_code == 403
    assert leader_client.get('/admin/zwroty/licencje').status_code == 302


def test_strony_dla_admina(admin_client):
    for adres in ('/admin/zwroty/licencje', '/admin/zwroty/cennik', '/admin/panel'):
        assert admin_client.get(adres).status_code == 200, adres
    assert 'Cennik zwrotów' in admin_client.get('/admin/panel').get_data(as_text=True)
