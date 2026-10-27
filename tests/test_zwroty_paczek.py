"""Zwroty Paczek (2026-10-26): naglowki 1:1 z arkusza operacji, filtry
(domyslnie dzis), eksport do Excela z tych samych wierszy co ekran."""
import io
import re
from datetime import date

import openpyxl

from conftest import login, make_user

from logistat.czas import local_today
from logistat.zwroty_dane import kurierzy_raportu, naglowki_zwrotow

# Arkusz operacji + kolumna „Other” (2026-10-27: kurierzy z Ustawien, kazdy osobna kolumna).
NAGLOWKI_OPERACJI = [
    'Paketnummer', 'Brand', 'Date', 'POS-Nummer', 'DPD', 'UPS', 'Other', 'Retoure', 'Defekt',
    'Neuwarenueberschuss', 'Silber_in_Stück', 'Defektes_Silber_in_Stück',
    'Reduzierte_ware', 'Verpackungsvorschrift', 'Monat', 'Woche',
]


def test_naglowki_jak_w_arkuszu_operacji(flask_app):
    dzis = local_today()
    assert list(naglowki_zwrotow(kurierzy_raportu(dzis, dzis))) == NAGLOWKI_OPERACJI


def test_ekran_ma_naglowki_i_domyslnie_dzis(leader_client):
    html = leader_client.get('/zwroty/paczki').get_data(as_text=True)

    naglowki = re.findall(r'<th>([^<]*)</th>', html)
    assert naglowki == NAGLOWKI_OPERACJI
    dzis = local_today().isoformat()
    assert f'name="date_from" value="{dzis}"' in html
    assert f'name="date_to" value="{dzis}"' in html
    assert '(dziś)' in html


def test_filtry_przechodza_do_linku_excela(leader_client):
    html = leader_client.get('/zwroty/paczki?date_from=2026-09-01&date_to=2026-09-30'
                             '&brand=ACC&paketnummer=123').get_data(as_text=True)

    link = re.search(r'href="(/zwroty/paczki/export[^"]*)"', html).group(1)
    assert 'date_from=2026-09-01' in link and 'date_to=2026-09-30' in link
    assert 'brand=ACC' in link and 'paketnummer=123' in link


def test_bledna_data_wraca_do_dzis(leader_client):
    r = leader_client.get('/zwroty/paczki?date_from=xx&date_to=')
    assert r.status_code == 200
    assert f'value="{local_today().isoformat()}"' in r.get_data(as_text=True)


def test_odwrocony_zakres_jest_zamieniany(leader_client):
    html = leader_client.get('/zwroty/paczki?date_from=2026-09-30&date_to=2026-09-01').get_data(as_text=True)
    assert 'name="date_from" value="2026-09-01"' in html


def test_excel_ma_naglowki_operacji(leader_client):
    r = leader_client.get('/zwroty/paczki/export?date_from=2026-09-02&date_to=2026-09-02')

    assert r.status_code == 200
    assert 'Returns_2026-09-02.xlsx' in r.headers['Content-Disposition']
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert [c.value for c in ws[1]] == NAGLOWKI_OPERACJI
    assert ws.auto_filter.ref


def test_excel_wiersze_z_tej_samej_funkcji_co_ekran(leader_client, monkeypatch):
    wiersz = ('%0000000000000000001', 'ACC', date(2026, 9, 2), 'NO', 0, 1, 0, 2, 1,
              0, 0, 0, 0, 'Yes', 'wrzesień-2026', '36-2026')
    wywolania = []

    def stub(**f):
        wywolania.append(f)
        return [wiersz]

    monkeypatch.setattr('logistat.web.zwroty.wiersze_zwrotow', stub)
    html = leader_client.get('/zwroty/paczki?date_from=2026-09-02').get_data(as_text=True)
    r = leader_client.get('/zwroty/paczki/export?date_from=2026-09-02')

    assert '%0000000000000000001' in html and '02.09.2026' in html
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert ws['A2'].value == '%0000000000000000001' and ws['A2'].number_format == '@'
    assert ws['C2'].value.date() == date(2026, 9, 2)
    assert wywolania[0] == wywolania[1] == {'od': date(2026, 9, 2), 'do': date(2026, 9, 2),
                                            'brand': None, 'paketnummer': None,
                                            'kurierzy': ['DPD', 'UPS', 'Other']}


def test_tylko_lider_i_admin(client):
    u = make_user('stanowisko')
    login(client, u.username)
    assert client.get('/zwroty/paczki').status_code == 302
    assert client.get('/zwroty/paczki/export').status_code == 302
