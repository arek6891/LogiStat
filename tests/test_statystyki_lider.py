"""Statystyki ogolne dla lidera (2026-10-12): podglad bez kosztow.

Lider widzi linie i ilosci, ale nie stawki ani koszty — ani na ekranie, ani
w zrodle strony (stawki szly do JS), ani w eksporcie Excel. Edycja ilosci
i „→ uzyj skanow" zostaja u admina; API zapisu dalej @admin_required.
"""
import io
from datetime import date

import openpyxl

import app as logistat

LADOWANIE = date(2026, 9, 10)
STRONA = f'/general-stats?date_from={LADOWANIE}&date_to={LADOWANIE}'
STAWKA = 7.7731   # charakterystyczna — latwo wypatrzyc w HTML


def linia():
    s = logistat.GeneralStat(loading_date=LADOWANIE, week_number=37, list_id='UB-1',
                             country_ledger='PL', country_of_destination='Polska',
                             amounts=100)
    s.set_category_data({'labelling_on': {'amount': 40, 'cost': 0.0}})
    mapping = logistat.CostMapping(year=LADOWANIE.year, month=LADOWANIE.month)
    mapping.set_rates_data({'labelling_on': STAWKA})
    logistat.db.session.add_all([s, mapping])
    logistat.db.session.commit()
    return s


def naglowki(ws):
    return [c.value for wiersz in ws.iter_rows(min_row=1, max_row=2) for c in wiersz
            if c.value is not None]


# ── ekran ────────────────────────────────────────────────────────────────────

def test_lider_widzi_ilosci_bez_kosztow(leader_client):
    linia()

    r = leader_client.get(STRONA)

    assert r.status_code == 200
    html = r.get_data(as_text=True)
    assert 'UB-1' in html and 'data-val="40"' in html
    assert 'Cost (EUR)' not in html
    assert 'Total cost' not in html
    assert 'class="total-cost"' not in html and 'id="sumTotalCost"' not in html
    assert 'cat-cost' not in html.split('<script>')[0]
    assert str(STAWKA) not in html, 'stawki nie moga trafic do zrodla strony'


def test_lider_nie_edytuje(leader_client):
    linia()

    html = leader_client.get(STRONA).get_data(as_text=True)

    assert 'class="cat-amount inline-input"' not in html
    assert 'readonly' in html
    assert 'autoZapisStatus"' not in html
    assert 'onclick="przelaczNaSkany' not in html


def test_admin_dalej_widzi_koszty_i_edytuje(admin_client):
    linia()

    html = admin_client.get(STRONA).get_data(as_text=True)

    assert 'Total cost (EUR)' in html
    assert 'Cost (EUR)' in html
    assert 'class="cat-amount inline-input"' in html


def test_lider_nie_zapisze_przez_api(leader_client):
    s = linia()

    r = leader_client.put(f'/api/general-stats/{s.id}',
                          json={'category_data': {'labelling_on': {'amount': 999}}})

    assert r.status_code != 200
    logistat.db.session.expire_all()
    assert logistat.db.session.get(logistat.GeneralStat, s.id) \
        .get_category_data()['labelling_on']['amount'] == 40
    assert leader_client.post(f'/api/general-stats/{s.id}/use-scan').status_code != 200
    assert leader_client.get('/api/general-stats').status_code != 200, \
        'API zwraca total_cost — zostaje dla admina'


# ── eksport ──────────────────────────────────────────────────────────────────

def test_eksport_lidera_bez_kolumn_kosztow(leader_client):
    linia()

    r = leader_client.get('/general-stats/export')

    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    h = naglowki(ws)
    assert 'Cost (EUR)' not in h and 'Total cost (EUR)' not in h
    # Jedna kolumna na kategorie: pierwsza kategoria zaraz po Total Amount.
    assert ws.cell(row=1, column=9).value == \
        logistat.STAT_CATEGORY_LABELS[logistat.STAT_CATEGORIES[0]]
    assert ws.max_column == 8 + len(logistat.STAT_CATEGORIES)
    ilosci = [ws.cell(row=3, column=9 + i).value for i in range(len(logistat.STAT_CATEGORIES))]
    assert ilosci[logistat.STAT_CATEGORIES.index('labelling_on')] == 40


def test_eksport_admina_z_kosztami(admin_client):
    linia()

    ws = openpyxl.load_workbook(
        io.BytesIO(admin_client.get('/general-stats/export').data)).active

    h = naglowki(ws)
    assert 'Cost (EUR)' in h and 'Total cost (EUR)' in h
    assert ws.max_column == 8 + 2 * len(logistat.STAT_CATEGORIES) + 1


def test_filtr_po_lewej_nad_tabela(leader_client):
    """Filtr byl po prawej obok Excela i nikt go nie widzial — teraz pod tytulem."""
    html = leader_client.get('/general-stats').get_data(as_text=True)

    assert 'Loading date od' in html and 'Loading date do' in html
    assert html.index('id="filtrForm"') < html.index('id="exportBtn"') < html.index('id="statsTable"')
