"""Dekoracje — dostawa produktow akcji (operacja 2026-10-29).

- typ dostawy (kontenerowa / Beeline) wymagany przy tworzeniu akcji, zmiana tylko
  dopoki akcja nie ma wpisow (409),
- kontenerowa: stanowisko wpisuje nr kontenera, potem wiersze artykul / kartony /
  PU; numery znormalizowane (spacje, wielkie litery) → sumy per kontener i per
  artykul; ten sam `token` drugi raz nie dodaje wiersza,
- Beeline: skan kartonu, ten sam karton w akcji drugi raz → 409 (kto / kiedy),
- zly typ → 400, zakonczona akcja → 409, akcja bez typu → 400,
- lider: ✕ wiersza / kartonu, Excel (numery jako tekst); stanowisko eksportu nie ma.
Dane syntetyczne.
"""
import io

import openpyxl
import pytest
from conftest import login, make_user
from test_dekoracje import nowa_akcja

from logistat.dekoracje import sumy_dostaw
from logistat.extensions import db
from logistat.models import AkcjaDekoracji, KartonBeeline, KontenerDostawy, PozycjaDostawy

_nr = iter(range(1, 10_000))


@pytest.fixture
def stanowisko(flask_app):
    klient = flask_app.test_client()
    login(klient, make_user('stanowisko').username)
    return klient


def wiersz(klient, akcja, kontener, artykul, kartony, pu, status=200, token=None):
    token = token or f'token-{next(_nr):06d}'
    r = klient.post(f'/api/dekoracje/akcje/{akcja["id"]}/kontener',
                    json={'kontener': kontener, 'artykul': artykul, 'kartony': kartony, 'pu': pu, 'token': token})
    assert r.status_code == status, r.get_json()
    return r.get_json()


def karton(klient, akcja, numer, status=201):
    r = klient.post(f'/api/dekoracje/akcje/{akcja["id"]}/beeline', json={'numer': numer})
    assert r.status_code == status, r.get_json()
    return r.get_json()


# ── czysta logika ────────────────────────────────────────────────────────────

def test_sumy_dostaw():
    kontenery, artykuly, razem = sumy_dostaw([
        {'kontener': 'K2', 'artykul': 'B', 'kartony': 1, 'pu': 10},
        {'kontener': 'K1', 'artykul': 'A', 'kartony': 2, 'pu': 0},
        {'kontener': 'K2', 'artykul': 'A', 'kartony': 3, 'pu': 5},
    ])
    assert [(k['kontener'], k['wierszy'], k['kartony'], k['pu']) for k in kontenery] == [('K2', 2, 4, 15), ('K1', 1, 2, 0)]
    assert [(a['artykul'], a['kartony'], a['pu']) for a in artykuly] == [('A', 5, 5), ('B', 1, 10)]
    assert razem == {'wierszy': 3, 'kartony': 6, 'pu': 15}


# ── typ dostawy ──────────────────────────────────────────────────────────────

def test_typ_wymagany_przy_tworzeniu(leader_client):
    for typ in (None, '', 'statek'):
        r = leader_client.post('/api/dekoracje/akcje', json={'nazwa': 'X', 'od': '2026-11-01', 'do': '2026-11-02',
                                                              'typ_dostawy': typ})
        assert r.status_code == 400 and 'typ dostawy' in r.get_json()['error']
    a = nowa_akcja(leader_client, typ_dostawy='beeline')
    assert a['typ_dostawy'] == 'beeline'


def test_zmiana_typu_tylko_bez_wpisow(leader_client, stanowisko):
    a = nowa_akcja(leader_client, typ_dostawy='beeline')
    url = f'/api/dekoracje/akcje/{a["id"]}/typ-dostawy'
    r = leader_client.put(url, json={'typ_dostawy': 'kontener'})
    assert r.status_code == 200 and r.get_json()['typ'] == 'kontener' and not r.get_json()['ma_wpisy']
    wiersz(stanowisko, a, 'K1', 'A1', 1, 1)
    r = leader_client.put(url, json={'typ_dostawy': 'beeline'})
    assert r.status_code == 409 and 'wpisy' in r.get_json()['error']
    assert leader_client.put(url, json={'typ_dostawy': 'kontener'}).status_code == 200, 'ten sam typ — bez zmian'
    assert leader_client.put(url, json={'typ_dostawy': 'x'}).status_code == 400


def test_akcja_bez_typu_nie_przyjmuje_wpisow(leader_client, stanowisko):
    """Akcja sprzed typow dostawy (kolumna NULL po migracji)."""
    a = nowa_akcja(leader_client)
    db.session.get(AkcjaDekoracji, a['id']).typ_dostawy = None
    db.session.commit()
    r = stanowisko.post(f'/api/dekoracje/akcje/{a["id"]}/kontener',
                        json={'kontener': 'K1', 'artykul': 'A', 'kartony': 1, 'pu': 1, 'token': 'token-bez-typu'})
    assert r.status_code == 400 and 'nie ustawił' in r.get_json()['error']
    strona = stanowisko.get('/dekoracje/dostawy').get_data(as_text=True)
    assert '"typ": null' in strona
    d = leader_client.get(f'/displays/akcje/{a["id"]}').get_data(as_text=True)
    assert 'wybierz' in d


# ── kontenerowa ──────────────────────────────────────────────────────────────

def test_kontener_sumy_i_normalizacja(leader_client, stanowisko):
    a = nowa_akcja(leader_client)
    r = stanowisko.get(f'/api/dekoracje/akcje/{a["id"]}/kontener?kod=msku 123')
    assert r.status_code == 200 and r.get_json() == {
        'kontener': 'MSKU123', 'nowy': True, 'wiersze': [], 'razem': {'wierszy': 0, 'kartony': 0, 'pu': 0}}
    wiersz(stanowisko, a, 'MSKU123', 'art-1', 3, 120)
    k = wiersz(stanowisko, a, 'msku123 ', ' ART-1', 2, 80)
    assert not k['nowy'] and k['razem'] == {'wierszy': 2, 'kartony': 5, 'pu': 200}
    wiersz(stanowisko, a, 'TGHU9', 'ART-2', 1, 0)
    assert KontenerDostawy.query.count() == 2

    d = leader_client.get(f'/api/dekoracje/akcje/{a["id"]}/dostawy').get_json()
    assert [(k['kontener'], k['kartony'], k['pu']) for k in d['kontenery']] == [('MSKU123', 5, 200), ('TGHU9', 1, 0)]
    assert [(x['artykul'], x['kartony'], x['pu']) for x in d['artykuly']] == [('ART-1', 5, 200), ('ART-2', 1, 0)]
    assert d['razem'] == {'wierszy': 3, 'kartony': 6, 'pu': 200} and d['ma_wpisy']


def test_kontener_ten_sam_token_raz(leader_client, stanowisko):
    a = nowa_akcja(leader_client)
    wiersz(stanowisko, a, 'K1', 'A', 1, 1, token='podwojny-enter')
    k = wiersz(stanowisko, a, 'K1', 'A', 1, 1, token='podwojny-enter')
    assert k['razem']['wierszy'] == 1 and PozycjaDostawy.query.count() == 1


@pytest.mark.parametrize('pola, blad', [
    ({'kartony': 0, 'pu': 0}, 'kartonów albo PU'),
    ({'kartony': -1}, 'Kartony'),
    ({'pu': 100_000}, 'PU'),
    ({'kartony': '3'}, 'Kartony'),
    ({'artykul': '  '}, 'artykułu'),
    ({'kontener': ''}, 'kontenera'),
    ({'token': ''}, 'identyfikatora'),
])
def test_kontener_bledne_dane(leader_client, stanowisko, pola, blad):
    a = nowa_akcja(leader_client)
    body = {'kontener': 'K1', 'artykul': 'A', 'kartony': 1, 'pu': 1, 'token': 'token-bledny', **pola}
    r = stanowisko.post(f'/api/dekoracje/akcje/{a["id"]}/kontener', json=body)
    assert r.status_code == 400 and blad in r.get_json()['error']
    assert PozycjaDostawy.query.count() == 0


def test_zly_typ_i_zakonczona_akcja(leader_client, stanowisko):
    a = nowa_akcja(leader_client, typ_dostawy='kontener')
    karton(stanowisko, a, 'BEE-1', status=400)
    b = nowa_akcja(leader_client, nazwa='Beeline', typ_dostawy='beeline')
    wiersz(stanowisko, b, 'K1', 'A', 1, 1, status=400)
    leader_client.put(f'/api/dekoracje/akcje/{a["id"]}/status', json={'zakonczona': True})
    r = stanowisko.post(f'/api/dekoracje/akcje/{a["id"]}/kontener',
                        json={'kontener': 'K1', 'artykul': 'A', 'kartony': 1, 'pu': 1, 'token': 'token-zakonczona'})
    assert r.status_code == 409
    assert leader_client.put(f'/api/dekoracje/akcje/{a["id"]}/typ-dostawy',
                             json={'typ_dostawy': 'beeline'}).status_code == 409
    strona = stanowisko.get('/dekoracje/dostawy').get_data(as_text=True)
    assert '"nazwa": "Beeline"' in strona and 'Akcja jesie' not in strona, 'na stanowisku tylko akcje w toku'


def test_usuniecie_ostatniego_wiersza_usuwa_kontener(leader_client, stanowisko):
    a = nowa_akcja(leader_client)
    wiersz(stanowisko, a, 'K1', 'A', 1, 1)
    wiersz(stanowisko, a, 'K1', 'B', 2, 2)
    ids = [p.id for p in PozycjaDostawy.query.order_by(PozycjaDostawy.id)]
    assert stanowisko.delete(f'/api/dekoracje/dostawy/pozycje/{ids[0]}').status_code == 403
    d = leader_client.delete(f'/api/dekoracje/dostawy/pozycje/{ids[0]}').get_json()
    assert d['razem']['wierszy'] == 1 and KontenerDostawy.query.count() == 1
    d = leader_client.delete(f'/api/dekoracje/dostawy/pozycje/{ids[1]}').get_json()
    assert not d['ma_wpisy'] and KontenerDostawy.query.count() == 0


# ── Beeline ──────────────────────────────────────────────────────────────────

def test_beeline_duplikat_odrzucony(leader_client, stanowisko):
    a = nowa_akcja(leader_client, typ_dostawy='beeline')
    assert karton(stanowisko, a, 'bee 001')['numer'] == 'BEE001'
    assert karton(stanowisko, a, 'BEE002')['kartonow'] == 2
    r = stanowisko.post(f'/api/dekoracje/akcje/{a["id"]}/beeline', json={'numer': 'BEE001'})
    assert r.status_code == 409 and 'już zeskanowany' in r.get_json()['error']
    assert KartonBeeline.query.count() == 2
    # Ten sam karton w innej akcji — wolno.
    b = nowa_akcja(leader_client, nazwa='Druga', typ_dostawy='beeline')
    karton(stanowisko, b, 'BEE001')

    k = KartonBeeline.query.filter_by(akcja_id=a['id'], numer='BEE001').one()
    d = leader_client.delete(f'/api/dekoracje/dostawy/beeline/{k.id}').get_json()
    assert [x['numer'] for x in d['kartony']] == ['BEE002']
    karton(stanowisko, a, 'BEE001')


# ── eksport i uprawnienia ────────────────────────────────────────────────────

def _arkusze(odpowiedz):
    assert odpowiedz.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(odpowiedz.data))
    return {ws.title: ws for ws in wb.worksheets}


def test_eksport_kontenerow(leader_client, stanowisko):
    a = nowa_akcja(leader_client)
    wiersz(stanowisko, a, '0012', '00077', 3, 120)
    wiersz(stanowisko, a, '0012', '00077', 1, 40)
    ark = _arkusze(leader_client.get(f'/displays/akcje/{a["id"]}/dostawy/export'))
    assert list(ark) == ['Wiersze', 'Per kontener', 'Per artykuł']
    w = list(ark['Wiersze'].iter_rows(values_only=True))
    assert w[0] == ('Kontener', 'Artykuł', 'Kartony', 'PU', 'Wpisano', 'Kto')
    assert w[1][:4] == ('0012', '00077', 3, 120)
    assert ark['Wiersze'].cell(2, 2).number_format == '@'
    assert list(ark['Per artykuł'].iter_rows(values_only=True))[1] == ('00077', 4, 160)


def test_eksport_beeline_i_uprawnienia(leader_client, stanowisko):
    a = nowa_akcja(leader_client, typ_dostawy='beeline')
    karton(stanowisko, a, '000123')
    ark = _arkusze(leader_client.get(f'/displays/akcje/{a["id"]}/dostawy/export'))
    w = list(ark['Kartony Beeline'].iter_rows(values_only=True))
    assert w[0] == ('Nr kartonu', 'Zeskanowano', 'Kto') and w[1][0] == '000123'
    assert stanowisko.get(f'/displays/akcje/{a["id"]}/dostawy/export').status_code == 302
    assert stanowisko.get(f'/api/dekoracje/akcje/{a["id"]}/dostawy').status_code == 403
    assert stanowisko.put(f'/api/dekoracje/akcje/{a["id"]}/typ-dostawy',
                          json={'typ_dostawy': 'kontener'}).status_code == 403
    # Stanowisko widzi swoj ekran w menu.
    assert 'Dostawy' in stanowisko.get('/dekoracje/dostawy').get_data(as_text=True)


def test_strony(leader_client):
    a = nowa_akcja(leader_client)
    d = leader_client.get(f'/displays/akcje/{a["id"]}').get_data(as_text=True)
    assert '🚚 Dostawa produktów' in d and 'Dostawa kontenerowa' in d
    assert 'Dostawa z Beeline' in leader_client.get('/displays').get_data(as_text=True)
