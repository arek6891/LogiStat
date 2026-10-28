"""Dekoracje → Etykiety POS (operacja 2026-10-28).

- Code 128 (B / C) i PDF z arkuszem etykiet — bez bibliotek, sprawdzane tu,
- import listy z Excela: kolumna POS wymagana, Post Processing opcjonalna
  (Tak / Nie / puste), numer = klucz (ponowny import aktualizuje, nie dubluje),
  bledne wiersze wracaja z numerem wiersza, nic nie ginie po cichu,
- eksport → poprawka → import w tym samym ukladzie,
- druk: wszystkie / z ostatniego importu / zaznaczone, kopie,
- lider tak, konto stanowiska nie.
Dane syntetyczne.
"""
import io
import re

import openpyxl
import pytest
from conftest import login, make_user

from logistat.dekoracje import flaga_pp, wiersze_pos
from logistat.kod128 import CISZA, WZORY, moduly, wartosci
from logistat.models import PosDekoracji
from logistat.pdf_etykiet import NA_STRONIE, miesci_sie, pdf_etykiet


def xlsx(wiersze):
    wb = openpyxl.Workbook()
    ws = wb.active
    for w in wiersze:
        ws.append(list(w))
    bufor = io.BytesIO()
    wb.save(bufor)
    bufor.seek(0)
    return bufor


def importuj(client, wiersze, status=200):
    r = client.post('/api/dekoracje/pos/import', data={'file': (xlsx(wiersze), 'pos.xlsx')},
                    content_type='multipart/form-data')
    assert r.status_code == status, r.get_json()
    return r.get_json()


def flagi():
    return {p.numer: p.post_processing for p in PosDekoracji.query.all()}


def etykiety_w_pdf(dane):
    return len(re.findall(rb'\) Tj ET', dane))


# ── kod kreskowy i PDF ───────────────────────────────────────────────────────

def test_wzory_kodu_128():
    assert len(WZORY) == 106 and len(set(WZORY)) == 106
    assert all(sum(map(int, w)) == 11 for w in WZORY)


def test_suma_kontrolna_i_zestaw_c():
    # Start B = 104; (104 + 33·1 + 34·2 + 35·3) mod 103 = 1
    assert wartosci('ABC') == [104, 33, 34, 35, 1]
    assert wartosci('123456')[:4] == [105, 12, 34, 56], 'same cyfry, parzyscie → zestaw C'
    assert wartosci('12345')[0] == 104, 'nieparzysta liczba cyfr → zestaw B'
    kod = moduly('ABC')
    assert kod.startswith('0' * CISZA) and kod.endswith('11' + '0' * CISZA)
    assert len(kod) == 2 * CISZA + 11 * 5 + 13
    with pytest.raises(ValueError):
        wartosci('')
    with pytest.raises(ValueError):
        wartosci('ŻÓŁW')


def test_pdf_arkusza():
    numery = [f'PL-{i:04d}' for i in range(NA_STRONIE + 1)]
    pdf = pdf_etykiet(numery)
    assert pdf.startswith(b'%PDF-1.4') and pdf.rstrip().endswith(b'%%EOF')
    assert b'/Count 2' in pdf, 'druga strona dla etykiety ponad arkusz'
    assert etykiety_w_pdf(pdf) == len(numery)
    # Tablica xref wskazuje poczatki obiektow — inaczej czytnik PDF „naprawia” plik.
    xref = int(pdf.rsplit(b'startxref\n', 1)[1].split(b'\n')[0])
    wpisy = pdf[xref:].split(b'\n')[3:]
    for nr, wpis in enumerate(wpisy[:6], start=1):
        przes = int(wpis[:10])
        assert pdf[przes:].startswith(f'{nr} 0 obj'.encode())
    with pytest.raises(ValueError):
        pdf_etykiet([])


def test_za_dlugi_numer_nie_miesci_sie():
    assert miesci_sie('POS-1234567890')
    assert not miesci_sie('A' * 30)


# ── parsowanie arkusza ───────────────────────────────────────────────────────

@pytest.mark.parametrize('komorka,wynik', [
    (None, (None, True)), ('  ', (None, True)), ('Tak', (True, True)), ('nie', (False, True)),
    ('YES', (True, True)), ('x', (True, True)), (1, (True, True)), (0, (False, True)),
    (True, (True, True)), ('może', (None, False)), (5, (None, False)),
])
def test_flaga_pp(komorka, wynik):
    assert flaga_pp(komorka) == wynik


def test_wiersze_pos():
    pozycje, ma_pp, bledy = wiersze_pos([
        ('Lp', '  Nr  POS ', 'Post-Processing'),
        (1, 'pl-100', 'Tak'),
        (2, 12345.0, None),                 # liczba z Excela — bez „.0”
        (3, None, None),                    # pusty wiersz — pomijany
        (4, ' pl-100 ', 'Nie'),             # ten sam numer — wygrywa ostatni
        (5, 'ZŁY#1', 'Tak'),
        (6, 'PL-7', 'może'),
        (7, None, 'Tak'),
        (8, 'A' * 30, None),
    ])
    assert ma_pp
    assert pozycje == {'PL-100': False, '12345': None}
    assert [(b['wiersz'], b['powod'].split()[0]) for b in bledy] == [
        (6, 'niedozwolone'), (7, 'Post'), (8, 'brak'), (9, 'za')]


def test_wiersze_pos_bez_kolumny():
    with pytest.raises(ValueError):
        wiersze_pos([('Numer', 'X'), ('PL-1', 1)])
    pozycje, ma_pp, _ = wiersze_pos([('POS',), ('PL-1',)])
    assert pozycje == {'PL-1': None} and not ma_pp


# ── import / eksport ─────────────────────────────────────────────────────────

def test_import_i_ponowny_import(leader_client):
    w = importuj(leader_client, [('POS', 'Post Processing'), ('PL-1', 'Tak'), ('PL-2', None), ('PL-3', 'Nie')])
    assert (w['dodane'], w['zmienione'], w['wczytane'], w['ostatni_import']) == (3, 0, 3, 3)
    assert flagi() == {'PL-1': True, 'PL-2': None, 'PL-3': False}

    w = importuj(leader_client, [('POS', 'Post Processing'), ('PL-2', 'Tak'), ('PL-3', 'Nie'), ('PL-4', ''), ('$$', '')])
    assert (w['dodane'], w['zmienione'], w['bez_zmian'], len(w['bledy'])) == (1, 1, 1, 1)
    assert w['ostatni_import'] == 3 and len(w['lista']) == 4
    assert flagi() == {'PL-1': True, 'PL-2': True, 'PL-3': False, 'PL-4': None}

    w = importuj(leader_client, [('POS',), ('PL-1',), ('PL-2',)])
    assert not w['ma_pp'] and w['zmienione'] == 0
    assert flagi()['PL-1'] is True, 'plik bez kolumny Post Processing jej nie zeruje'


def test_import_bledny_plik(leader_client):
    r = leader_client.post('/api/dekoracje/pos/import', data={'file': (io.BytesIO(b'nie excel'), 'a.xlsx')},
                           content_type='multipart/form-data')
    assert r.status_code == 400
    assert 'POS' in importuj(leader_client, [('Numer',), ('PL-1',)], status=400)['error']
    importuj(leader_client, [('POS',)], status=400)
    assert leader_client.post('/api/dekoracje/pos/import').status_code == 400


def test_eksport_wczytuje_sie_z_powrotem(leader_client):
    importuj(leader_client, [('POS', 'Post Processing'), ('00123', 'Tak'), ('PL-2', None)])
    r = leader_client.get('/displays/pos/export')
    assert r.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    wiersze = list(ws.iter_rows(values_only=True))
    assert wiersze == [('POS', 'Post Processing'), ('00123', 'Tak'), ('PL-2', None)]
    assert ws.cell(2, 1).number_format == '@', 'numer jako tekst (wiodace zera)'
    w = importuj(leader_client, wiersze)
    assert (w['dodane'], w['zmienione'], w['bez_zmian']) == (0, 0, 2)


# ── druk ─────────────────────────────────────────────────────────────────────

def test_druk_zakresy_i_kopie(leader_client):
    importuj(leader_client, [('POS',), ('PL-1',), ('PL-2',), ('PL-3',)])
    importuj(leader_client, [('POS',), ('PL-3',), ('PL-4',)])
    r = leader_client.post('/displays/pos/etykiety', data={'zakres': 'wszystkie'})
    assert r.status_code == 200 and r.headers['Content-Type'] == 'application/pdf'
    assert r.headers['Content-Disposition'].startswith('inline')
    assert etykiety_w_pdf(r.data) == 4
    r = leader_client.post('/displays/pos/etykiety', data={'zakres': 'import', 'kopie': '2'})
    assert etykiety_w_pdf(r.data) == 4 and b'(PL-4) Tj' in r.data and b'(PL-1) Tj' not in r.data
    ids = [str(p.id) for p in PosDekoracji.query.filter(PosDekoracji.numer.in_(['PL-1', 'PL-2']))]
    r = leader_client.post('/displays/pos/etykiety', data={'zakres': 'wybrane', 'id': ids + ['x', '999999']})
    assert etykiety_w_pdf(r.data) == 2


@pytest.mark.parametrize('dane', [
    {'zakres': 'wybrane'}, {'zakres': 'wszystkie', 'kopie': '0'}, {'zakres': 'wszystkie', 'kopie': 'x'},
    {'zakres': 'cos'},
])
def test_druk_bez_etykiet_wraca_z_komunikatem(leader_client, dane):
    importuj(leader_client, [('POS',), ('PL-1',)])
    r = leader_client.post('/displays/pos/etykiety', data=dane)
    assert r.status_code == 302 and r.headers['Location'].endswith('/displays/pos')


def test_usun_i_strona(leader_client):
    importuj(leader_client, [('POS', 'PP'), ('PL-1', 'tak'), ('PL-2', 'nie')])
    p = PosDekoracji.query.filter_by(numer='PL-1').one()
    assert leader_client.delete(f'/api/dekoracje/pos/{p.id}').status_code == 200
    assert leader_client.delete(f'/api/dekoracje/pos/{p.id}').status_code == 404
    html = leader_client.get('/displays/pos').get_data(as_text=True)
    assert '🏷️ Etykiety POS' in html and 'PL-2' in html and 'PL-1' not in html


def test_stanowisko_nie_ma_dostepu(flask_app):
    klient = flask_app.test_client()
    login(klient, make_user('stanowisko').username)
    r = klient.post('/api/dekoracje/pos/import', data={'file': (xlsx([('POS',), ('PL-1',)]), 'a.xlsx')},
                    content_type='multipart/form-data')
    assert r.status_code == 403
    r = klient.get('/displays/pos')
    assert r.status_code == 302 and '/start' in r.headers['Location']
