"""Total Amount w Statystykach ogolnych = wylacznie „Labelling one".

Suma wszystkich kategorii liczyla te same sztuki wielokrotnie — jedna sztuka
przechodzi przez kilka czynnosci, wiec „total" wychodzil wielokrotnoscia
zawartosci paczki i przebijal Amounts. Zolty wiersz double rate ZOSTAJE na
sumie wszystkich kategorii — decyzja do ustalenia z operacja (docs/TODO.md).
"""
import io
from datetime import date

import openpyxl

import app as logistat

LADOWANIE = date(2026, 9, 10)
ILOSCI = {
    'labelling_on':  40,
    'labelling_tvl': 7,
    'textile':       3,
    'sorting':       5,
}
KOLUMNA_TOTAL_AMOUNT = 8   # 'Loading date','Double Rate','Week','List-ID',
                           # 'Country of destination','Country ledger','Amounts','Total Amount'


def linia(list_id='UB-1', ilosci=None, amounts=100):
    s = logistat.GeneralStat(loading_date=LADOWANIE, week_number=37, list_id=list_id,
                             country_ledger='PL', country_of_destination='Polska',
                             amounts=amounts)
    s.set_category_data({k: {'amount': v, 'cost': 0.0}
                         for k, v in (ilosci if ilosci is not None else ILOSCI).items()})
    logistat.db.session.add(s)
    logistat.db.session.commit()
    return s


def eksport(admin_client):
    r = admin_client.get('/general-stats/export')
    assert r.status_code == 200
    return openpyxl.load_workbook(io.BytesIO(r.data)).active


# ── stala ────────────────────────────────────────────────────────────────────

def test_kategoria_total_amount_jest_w_liscie_kategorii():
    assert logistat.TOTAL_AMOUNT_CATEGORY in logistat.STAT_CATEGORIES


# ── eksport Excel ────────────────────────────────────────────────────────────

def test_eksport_liczy_total_amount_tylko_z_labelling_one(admin_client):
    linia()

    ws = eksport(admin_client)

    assert ws.cell(row=3, column=KOLUMNA_TOTAL_AMOUNT).value == 40, \
        'suma wszystkich kategorii dalaby 55'


def test_eksport_total_amount_zero_gdy_brak_labelling_one(admin_client):
    linia(ilosci={'sorting': 12, 'textile': 8})

    ws = eksport(admin_client)

    assert ws.cell(row=3, column=KOLUMNA_TOTAL_AMOUNT).value == 0


def test_zolty_wiersz_double_rate_zostaje_na_sumie(admin_client, acting_admin):
    """Swiadoma niespojnosc — do decyzji z operacja, patrz docs/TODO.md."""
    logistat.process_import_rows([{
        'barcode': 'B-DR', 'land': 'PL', 'stueckzahl': 10, 'kategorie': 'textile',
        'ziel_datum': LADOWANIE, 'uebergabe_nr': 'UB-1', 'double_rate': True,
    }])
    s = logistat.GeneralStat.query.filter_by(list_id='UB-1').first()
    s.set_category_data({k: {'amount': v, 'cost': 0.0} for k, v in ILOSCI.items()})
    s.set_double_rate_category_data({k: {'amount': v, 'cost': 0.0} for k, v in ILOSCI.items()})
    logistat.db.session.commit()

    ws = eksport(admin_client)

    assert ws.cell(row=3, column=KOLUMNA_TOTAL_AMOUNT).value == 40, 'zwykly wiersz'
    assert ws.cell(row=4, column=2).value == 'DOUBLE RATE'
    assert ws.cell(row=4, column=KOLUMNA_TOTAL_AMOUNT).value == 55, 'zolty wiersz: suma'


# ── ekran ────────────────────────────────────────────────────────────────────

def test_strona_pokazuje_total_amount_z_labelling_one(admin_client):
    linia()

    html = admin_client.get(
        f'/general-stats?date_from={LADOWANIE}&date_to={LADOWANIE}').get_data(as_text=True)

    assert 'class="row-total-amount"' in html
    assert 'data-val="40"' in html
    assert 'data-val="55"' not in html


def test_strona_ma_dymek_wyjasniajacy_total_amount(admin_client):
    linia()

    html = admin_client.get(
        f'/general-stats?date_from={LADOWANIE}&date_to={LADOWANIE}').get_data(as_text=True)

    assert 'podpowiedz-tresc' in html
    assert 'Labelling one' in html


# ── etykiety dwuczlonowe ─────────────────────────────────────────────────────

def test_etykieta_kategorii_jest_dwuczlonowa():
    assert logistat.etykieta_kategorii('labelling_on') == 'Labelling one — Etykietowanie pojedyncze'
    assert logistat.etykieta_kategorii('card_facture') == 'Card facture — Karta / faktura'


def test_kazda_kategoria_ma_polskie_tlumaczenie():
    brakujace = [k for k in logistat.STAT_CATEGORIES
                 if k not in logistat.STAT_CATEGORY_LABELS_PL]

    assert brakujace == []


def test_naglowki_eksportu_zostaja_po_angielsku(admin_client):
    """Eksport idzie na zewnatrz — polskie nazwy tylko na ekranach."""
    linia()

    ws = eksport(admin_client)
    naglowki = [c.value for c in ws[1] if c.value]

    assert 'Labelling one' in naglowki
    assert not any('Etykietowanie' in str(n) for n in naglowki)
