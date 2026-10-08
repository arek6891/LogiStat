"""Forecast: znak roznicy = Actual − Forecast.

Zgloszenie operacji (2026-10-08): plan 90 000, przyjechalo 20 000 — ekran
pokazywal +70 000 (Forecast − Actual), czytane jako nadwyzka. Minus ma znaczyc
„przyjechalo mniej niz plan" — w API, tabeli, podsumowaniu i eksporcie Excel.
"""
import io
from datetime import date

import openpyxl

import app as logistat

DZIEN = date(2026, 9, 10)


def przygotuj(client, plan, przyjechalo):
    assert client.post('/api/forecast/save', json={
        'date': DZIEN.isoformat(), 'quantity': plan}).status_code == 200
    logistat.db.session.add(logistat.ImportedCarton(
        barcode='FC-1', land='PL', stueckzahl=przyjechalo,
        data_pliku=DZIEN, uebergabe_nr='UB-1'))
    logistat.db.session.commit()


def test_mniej_niz_plan_to_minus(leader_client):
    przygotuj(leader_client, 90000, 20000)

    dane = leader_client.get(
        f'/api/forecast/chart-data?date_from={DZIEN}&date_to={DZIEN}').get_json()

    assert dane[0]['forecast'] == 90000 and dane[0]['actual'] == 20000
    assert dane[0]['diff'] == -70000


def test_wiecej_niz_plan_to_plus(leader_client):
    przygotuj(leader_client, 20000, 90000)

    dane = leader_client.get(
        f'/api/forecast/chart-data?date_from={DZIEN}&date_to={DZIEN}').get_json()

    assert dane[0]['diff'] == 70000


def test_eksport_excel_ma_ten_sam_znak(leader_client):
    przygotuj(leader_client, 90000, 20000)

    r = leader_client.get(f'/api/forecast/export?date_from={DZIEN}&date_to={DZIEN}')
    ws = openpyxl.load_workbook(io.BytesIO(r.data)).active
    naglowki = [c.value for c in ws[1]]
    kol = naglowki.index('Różnica (A-F)')
    wiersz = next(w for w in ws.iter_rows(min_row=2, values_only=True) if w[1] == 90000)

    assert wiersz[kol] == -70000
